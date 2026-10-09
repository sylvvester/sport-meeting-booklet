import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from openpyxl import load_workbook

from sport_meeting.core.booklet_renderer import render_competition_groups, render_full_booklet
from sport_meeting.core.heat_generation import generate_heat_assignments
from sport_meeting.core.importers import import_registration_workbooks, import_schedule_workbook
from sport_meeting.core.legacy_booklet_content import LAST_YEAR_SCHOOL_RECORDS
from sport_meeting.core.models import (
    Participant,
    PerformanceKind,
    Result,
    ResultStatus,
    Severity,
    ValidationIssue,
    ValidationReport,
    ProjectConfig,
    EventRound,
    EventType,
    RoundType,
)
from sport_meeting.core.reporting import (
    FinalMaterial,
    FinalRankingMaterial,
    render_final_assignments,
    render_final_assignments_batch,
    render_on_site_package,
    render_result_ranking,
    render_result_rankings_batch,
    render_validation_report,
)
from sport_meeting.core.workbook_templates import (
    create_grouping_registration_sample,
    create_grouping_schedule_sample,
)


class ReportingTests(unittest.TestCase):
    def test_booklet_schedule_groups_sessions_and_preserves_all_events(self):
        with tempfile.TemporaryDirectory() as directory:
            rounds = [
                EventRound("a", "50米", "初一男子组", EventType.TRACK, RoundType.PRELIMINARY, 4,
                           "2026-10-20 00:00 上午 09:30", 28),
                EventRound("b", "立定跳远", "初一女子组", EventType.FIELD, RoundType.TIMED_FINAL, 1,
                           "2026-10-20 00:00 上午 09:30", 25),
                EventRound("c", "4×100米接力", "高一女子组", EventType.TEAM, RoundType.TIMED_FINAL, 2,
                           "2026-10-21 00:00 下午 14:30", 12),
            ]
            output = Path(directory) / "booklet.docx"
            render_full_booklet(output, rounds, [], [])
            doc = Document(output)
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("10月20日（星期二） 上午", text)
            self.assertIn("10月21日（星期三） 下午", text)
            tables = [t for t in doc.tables if t.cell(0, 0).text == "比赛时间"]
            self.assertEqual(sum(len(t.rows) - 1 for t in tables), 3)
            self.assertEqual(tables[0].cell(1, 0).text, "09:30")
            self.assertEqual(tables[-1].cell(1, 3).text, "12队")
            for table in tables:
                self.assertEqual(len(table.columns), 5)
                self.assertIsNotNone(table.rows[0]._tr.trPr.find(qn("w:tblHeader")))
                self.assertGreater(table.columns[1].width, table.columns[0].width)
                self.assertNotIn("00:00", " ".join(c.text for row in table.rows for c in row.cells))

    def test_booklet_static_headings_notes_and_page_numbers(self):
        with tempfile.TemporaryDirectory() as directory:
            config = ProjectConfig(
                opening_ceremony_text="示例学校第十一届田径运动会开幕式议程\n主持人：甲\n一、方阵入场",
                officials_text="示例学校第十一届田径运动会 仲裁委员会及裁判员名单\n医务组：乙\n备注：人员可调整",
            )
            output = Path(directory) / "booklet.docx"
            render_full_booklet(output, [], [], [], config=config)
            doc = Document(output)
            texts = [p.text for p in doc.paragraphs]
            self.assertNotIn(config.opening_ceremony_text.splitlines()[0], texts)
            self.assertIn("主持人：甲", texts)
            label = next(p for p in doc.paragraphs if p.text == "主持人：甲")
            self.assertTrue(label.runs[0].bold)
            self.assertEqual(label.runs[0].font.size.pt, 12)
            preceding = doc.paragraphs[texts.index("备注：人员可调整") - 1]
            self.assertTrue(preceding.paragraph_format.keep_with_next)
            for p in doc.paragraphs:
                if p.style.name == "Heading 1":
                    self.assertEqual(str(p.runs[0].font.color.rgb), "000000")
                    self.assertTrue(p.paragraph_format.page_break_before)
            # 默认保留封面：封面不标页码，页码从目录页重新按 1 开始。
            self.assertTrue(doc.sections[0].different_first_page_header_footer)
            self.assertIn('w:instr="PAGE"', doc.sections[0].footer._element.xml)
            self.assertIsNone(doc.styles["Title"]._element.pPr.find(qn("w:pBdr")))

    def test_default_booklet_keeps_cover_without_page_number_and_restarts_at_contents(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "默认封面.docx"
            config = ProjectConfig(booklet_title="第十届测试秩序册")
            render_full_booklet(output, [], [], [], config=config)
            doc = Document(output)
            texts = [p.text for p in doc.paragraphs]
            # 封面保留，但封面所在节不标页码
            self.assertEqual(texts[0], config.booklet_title)
            self.assertIn("竞赛秩序册", texts)
            self.assertEqual(len(doc.sections), 2)
            cover_section = doc.sections[0]
            self.assertTrue(cover_section.different_first_page_header_footer)
            self.assertEqual(cover_section.first_page_footer.paragraphs[0].text.strip(), "")
            # 页码从目录所在节重新按 1 开始
            page_numbering = doc.sections[1]._sectPr.find(qn("w:pgNumType"))
            self.assertIsNotNone(page_numbering)
            self.assertEqual(page_numbering.get(qn("w:start")), "1")
            self.assertIn('w:instr="PAGE"', doc.sections[1].footer._element.xml)
            contents = next(p for p in doc.paragraphs if p.text == "目录")
            self.assertEqual(contents.style.name, "Heading 1")
            self.assertTrue(contents.paragraph_format.page_break_before)

    def test_booklet_without_cover_starts_at_table_of_contents(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "无封面.docx"
            config = ProjectConfig(booklet_title="第十届测试秩序册")
            render_full_booklet(output, [], [], [], config=config, include_cover=False)
            doc = Document(output)
            texts = [p.text for p in doc.paragraphs]
            self.assertEqual(texts[0], "目录")
            self.assertNotIn(config.booklet_title, texts)
            self.assertFalse(doc.sections[0].different_first_page_header_footer)
            self.assertIn('w:instr="PAGE"', doc.sections[0].footer._element.xml)

    def test_static_chapter_right_aligns_organizing_committee_signature(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "落款.docx"
            config = ProjectConfig(
                competition_rules_text="九、其它未尽事宜\n校运会组委会\n2026年10月",
            )
            render_full_booklet(output, [], [], [], config=config)
            doc = Document(output)
            committee = next(p for p in doc.paragraphs if p.text == "校运会组委会")
            date = next(p for p in doc.paragraphs if p.text == "2026年10月")
            self.assertEqual(committee.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
            self.assertEqual(date.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
            self.assertTrue(committee.paragraph_format.keep_with_next)

    def test_static_chapter_does_not_remove_body_mentioning_chapter_title(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "booklet.docx"
            render_full_booklet(output, [], [], [], config=ProjectConfig(
                opening_ceremony_text="请按开幕式议程提前到场\n时间：09:00",
            ))
            self.assertIn("请按开幕式议程提前到场", [p.text for p in Document(output).paragraphs])

    def test_booklet_schedule_orders_events_before_field_before_team(self):
        with tempfile.TemporaryDirectory() as directory:
            rounds = [
                EventRound("t1", "50米", "初一男子组", EventType.TRACK, RoundType.PRELIMINARY, 4,
                           "2026-10-20 00:00 上午 10:30", 28),
                EventRound("f1", "立定跳远", "初一女子组", EventType.FIELD, RoundType.TIMED_FINAL, 1,
                           "2026-10-20 00:00 上午 09:00", 25),
                EventRound("m1", "1分钟集体跳绳", "初一年级组", EventType.TEAM, RoundType.TIMED_FINAL, 1,
                           "2026-10-20 00:00 上午 08:30", 14),
            ]
            output = Path(directory) / "日程顺序.docx"
            render_full_booklet(output, rounds, [], [])
            doc = Document(output)
            tables = [t for t in doc.tables if t.cell(0, 0).text == "比赛时间"]
            self.assertEqual(len(tables), 3)
            # 顺序固定为径赛 → 田赛 → 集体项目，且各自按时间排列
            self.assertEqual(tables[0].cell(1, 1).text, "初一男子组50米")
            self.assertEqual(tables[1].cell(1, 1).text, "初一女子组立定跳远")
            self.assertEqual(tables[2].cell(1, 1).text, "初一年级组1分钟集体跳绳")

    def test_final_assignment_document(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "决赛.docx"
            participants = {
                "a": Participant("a", "10101", "甲", "初一(1)"),
                "b": Participant("b", "10201", "乙", "初一(2)"),
            }
            render_final_assignments("初一男子50米", "11:40", [[("a", 4), ("b", 5)]], participants, output)
            document = Document(output)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(document.tables[0].cell(1, 0).text, "4")
            self.assertEqual(round(document.sections[0].page_width.cm, 1), 21.0)

    def test_batch_final_document_uses_track_and_field_position_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "全部决赛.docx"
            participants = {
                "a": Participant("a", "10101", "甲", "初一(1)"),
                "b": Participant("b", "10201", "乙", "初一(2)"),
            }
            materials = [
                FinalMaterial("初一男子组50米", "11:40", [[("a", 4)]], "道次"),
                FinalMaterial("初一女子组跳远", "14:00", [[("b", 1)]], "出场序"),
            ]
            render_final_assignments_batch(materials, participants, output)
            document = Document(output)
            self.assertEqual(len(document.tables), 2)
            self.assertEqual(document.tables[0].cell(0, 0).text, "道次")
            self.assertEqual(document.tables[1].cell(0, 0).text, "出场序")

    def test_on_site_package_contains_four_ready_to_use_files(self):
        with tempfile.TemporaryDirectory() as directory:
            participants = {
                "a": Participant("a", "10101", "甲同学", "初一(1)"),
                "b": Participant("b", "10201", "乙同学", "初一(2)"),
            }
            materials = [
                FinalMaterial("初一男子50米", "2026-10-15 上午 11:40", [[("a", 4), ("b", 5)]]),
                FinalMaterial("初一女子跳远", "2026-10-15 下午 14:30", [[("b", 1)]], "出场序"),
            ]
            files = render_on_site_package(materials, participants, directory)
            for path in (
                files.assignment_sheet,
                files.check_in_sheet,
                files.official_sheet,
                files.broadcast_script,
            ):
                self.assertTrue(path.is_file())

            check_in = Document(files.check_in_sheet)
            self.assertEqual(check_in.tables[0].cell(0, 4).text, "到检")
            self.assertEqual(check_in.tables[0].cell(1, 4).text, "□")
            official = Document(files.official_sheet)
            official_text = "\n".join(paragraph.text for paragraph in official.paragraphs)
            self.assertIn("发令员道次表", official_text)
            self.assertIn("裁判员出场顺序表", official_text)
            broadcast = files.broadcast_script.read_text(encoding="utf-8-sig")
            self.assertIn("第4道，10101号，甲同学，初一(1)班", broadcast)
            self.assertIn("请以上运动员提前到检录处检录", broadcast)

    def test_ranking_uses_supplied_competition_rank(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "排名.docx"
            participant = Participant("a", "10101", "甲", "初一(1)")
            result = Result("a", 1, "12.34", PerformanceKind.TIME, 12340, "MILLISECOND", "12.34")
            render_result_ranking("初一男子50米", [(1, participant, result)], output)
            self.assertEqual(Document(output).tables[0].cell(1, 0).text, "1")

    def test_batch_ranking_document_keeps_blank_rank_for_invalid_result(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "全部排名.docx"
            participant = Participant("a", "10101", "甲", "初一(1)")
            result = Result(
                "a", 1, "DQ", PerformanceKind.TIME, None, None, "DQ",
                status=ResultStatus.DQ,
            )
            render_result_rankings_batch(
                [FinalRankingMaterial("初一男子50米", [(None, participant, result)])], output
            )
            document = Document(output)
            self.assertEqual(document.tables[0].cell(1, 0).text, "")

    def test_xlsx_report_uses_validation_report(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "核对.xlsx"
            issue = ValidationIssue(
                "IMPORT_999", Severity.WARNING, "测试错误", source_file_id="source-1",
                source_sheet="男生", source_row=17,
            )
            report = ValidationReport([issue], {"source-1": "示例选手73报名表.xlsx"})
            render_validation_report(report, output)
            workbook = load_workbook(output, read_only=True)
            self.assertEqual(workbook.active["A2"].value, "IMPORT_999")
            self.assertEqual(workbook.active["C1"].value, "具体问题")
            self.assertEqual(workbook.active["F2"].value, "示例选手73报名表.xlsx")
            self.assertEqual(workbook.active["G2"].value, "男生")
            self.assertEqual(workbook.active["H2"].value, 17)
            workbook.close()

    def test_competition_group_document_matches_five_column_layout(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            registration = import_registration_workbooks([
                create_grouping_registration_sample(root / "报名.xlsx")
            ])
            schedule = import_schedule_workbook(create_grouping_schedule_sample(root / "日程.xlsx"))
            generated = generate_heat_assignments(
                ProjectConfig(), registration.athletes, registration.entries, schedule.rounds, seed=2026
            )
            output = root / "竞赛分组表.docx"
            stats = render_competition_groups(
                output, schedule.rounds, generated.assignments, registration.athletes
            )
            document = Document(output)
            self.assertEqual(stats.table_count, 4)
            self.assertEqual(len(document.tables), 4)
            self.assertTrue(all(len(table.columns) == 5 for table in document.tables))
            self.assertTrue(all(
                cell.paragraphs[0].paragraph_format.keep_with_next
                for table in document.tables if len(table.rows) <= 12
                for row in table.rows[:-1] for cell in row.cells
            ))
            first_headers = [table.cell(0, 0).text for table in document.tables]
            self.assertEqual(first_headers.count("道次"), 3)
            self.assertEqual(first_headers.count("序号"), 1)
            self.assertEqual(round(document.sections[0].page_width.cm, 1), 21.0)
            self.assertEqual(round(document.sections[0].page_height.cm, 1), 29.7)

    def test_competition_group_chapter_labels_track_and_field_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            registration = import_registration_workbooks([
                create_grouping_registration_sample(root / "报名.xlsx")
            ])
            schedule = import_schedule_workbook(create_grouping_schedule_sample(root / "日程.xlsx"))
            generated = generate_heat_assignments(
                ProjectConfig(), registration.athletes, registration.entries, schedule.rounds, seed=2026
            )
            output = root / "竞赛分组表.docx"
            render_competition_groups(
                output, schedule.rounds, generated.assignments, registration.athletes
            )
            document = Document(output)
            # 分组表必须在“日期 时段”之下给出“径赛/田赛/集体项目”副标题
            labels = [
                p.text.strip() for p in document.paragraphs
                if p.text.strip() in ("径赛", "田赛", "集体项目")
            ]
            self.assertIn("径赛", labels)
            self.assertIn("田赛", labels)
            # 每个时段内副标题固定为 径赛 → 田赛 → 集体项目
            self.assertEqual(sorted(set(labels)), labels[: len(set(labels))])

    def test_full_booklet_contains_eight_chapters_and_dynamic_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            registration = import_registration_workbooks([
                create_grouping_registration_sample(root / "报名.xlsx")
            ])
            schedule = import_schedule_workbook(create_grouping_schedule_sample(root / "日程.xlsx"))
            generated = generate_heat_assignments(
                ProjectConfig(), registration.athletes, registration.entries, schedule.rounds, seed=2026
            )
            output = root / "完整秩序册.docx"
            config = ProjectConfig(
                booklet_title="第十届测试运动会秩序册",
                opening_ceremony_text="时间：上午八点\n地点：田径场",
                organizing_committee_text="主任：测试老师",
                officials_text="总裁判长：测试老师",
                competition_rules_text="一、每人最多报名两个项目。",
                school_records_text="男子100米  11.00秒",
            )
            stats = render_full_booklet(
                output, schedule.rounds, generated.assignments, registration.athletes,
                config=config, include_cover=True,
            )
            document = Document(output)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("第十届测试运动会秩序册", text)
            self.assertIn("地点：田径场", text)
            self.assertIn("主任：测试老师", text)
            self.assertIn("男子100米  11.00秒", text)
            for chapter in (
                "一、开幕式议程", "二、组委会名单", "三、仲裁委员会及裁判员名单",
                "四、竞赛规程", "五、竞赛日程", "六、运动员代表队名单",
                "七、竞赛分组表", "八、示例学校田径运动会校记录",
            ):
                self.assertIn(chapter, text)
            self.assertIn("初一(1)班", text)
            self.assertEqual(stats.table_count, 4)
            self.assertGreater(len(document.tables), stats.table_count)
            self.assertEqual(document.tables[0].cell(0, 0).text, "比赛时间")
            self.assertEqual(round(document.sections[0].page_width.cm, 1), 21.0)
            roster_tables = [table for table in document.tables if len(table.columns) == 6]
            self.assertTrue(roster_tables)
            self.assertTrue(all(
                cell.paragraphs[0].paragraph_format.keep_with_next
                for table in roster_tables if len(table.rows) <= 12
                for row in table.rows[:-1] for cell in row.cells
            ))

    def test_default_booklet_uses_last_year_static_content_and_latest_record_table(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "去年内容.docx"
            stats = render_full_booklet(output, [], [], [], config=ProjectConfig())
            document = Document(output)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("主持人：示例乙副校长", text)
            self.assertIn("组委会主任：示例甲", text)
            self.assertEqual(document.tables[-1].cell(0, 0).text, "项目")
            self.assertEqual(document.tables[-1].cell(1, 2).text, "15″10")

    def test_old_record_table_still_fits_long_record_holder(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "旧校纪录排版.docx"
            stats = render_full_booklet(
                output, [], [], [],
                config=ProjectConfig(school_records_text=LAST_YEAR_SCHOOL_RECORDS),
            )
            document = Document(output)
            record_holder_cell = next(
                row.cells[3]
                for row in document.tables[-1].rows
                if row.cells[3].text == "示例丙、示例丁"
            )
            run = next(run for run in record_holder_cell.paragraphs[0].runs if run.text)
            self.assertLess(run.font.size.pt, 10.5)
            self.assertIsNotNone(record_holder_cell._tc.tcPr.find(qn("w:noWrap")))
            self.assertGreater(stats.font_adjustment_count, 0)
            self.assertEqual(stats.layout_warnings, ())

    def test_full_booklet_renders_saved_static_content(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "静态章节.docx"
            config = ProjectConfig(
                booklet_title="第十届测试秩序册",
                opening_ceremony_text="时间：上午八点\n一、运动员入场",
                organizing_committee_text="主任：甲",
                officials_text="总裁判长：乙",
                competition_rules_text="一、按时检录",
                school_records_text="男子100米：11.80秒",
            )
            render_full_booklet(output, [], [], [], config=config, include_cover=True)
            text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
            for expected in (
                config.booklet_title, "时间：上午八点", "主任：甲", "总裁判长：乙",
                "一、按时检录", "男子100米：11.80秒",
            ):
                self.assertIn(expected, text)

    def test_full_booklet_reports_text_that_cannot_safely_fit_one_line(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "超长内容.docx"
            config = ProjectConfig(
                school_records_text=(
                    "项目\t组别\t成绩\t创造者\t运动会名称\n"
                    "引体向上\t高一男子\t23个\t这是一段明显超过单元格容量的超长姓名\t第九届校运会"
                )
            )
            stats = render_full_booklet(output, [], [], [], config=config)
            self.assertEqual(len(stats.layout_warnings), 1)
            self.assertIn("超长姓名", stats.layout_warnings[0])
            document = Document(output)
            cell = document.tables[-1].cell(1, 3)
            run = next(run for run in cell.paragraphs[0].runs if run.text)
            self.assertEqual(run.font.size.pt, 7.0)
            self.assertIsNone(cell._tc.tcPr.find(qn("w:noWrap")))


if __name__ == "__main__":
    unittest.main()
