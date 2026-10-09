import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from sport_meeting.core.database import ProjectDatabase
from sport_meeting.core.finals_workflow import compute_track_final
from sport_meeting.core.importers import import_registration_workbooks, import_schedule_workbook, parse_available_lanes
from sport_meeting.core.models import (
    ROUND_MODE_PRELIMINARY_FINAL,
    EventType,
    HeatAssignment,
    PerformanceKind,
    ProjectConfig,
    Result,
    RoundType,
    Severity,
)
from sport_meeting.core.workbook_templates import (
    SCHEDULE_HEADERS,
    REGISTRATION_HEADERS,
    create_registration_template,
    create_school_level_registration_template,
    create_schedule_template,
    create_schedule_test_sample,
    create_grouping_registration_sample,
    create_grouping_schedule_sample,
)
from sport_meeting.core.validation import validate_project
from sport_meeting.core.heat_generation import generate_heat_assignments


class ImportTests(unittest.TestCase):
    def test_detects_shifted_header_and_normalizes_event(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "报名.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "男生"
            sheet.append(["某届运动会报名表"])
            sheet.append(["说明"])
            sheet.append(["男生"])
            sheet.append(["男生项目"])
            sheet.append(["号码", "姓名", "性别", "年级", "班级", "项目1", "项目2"])
            sheet.append([10901, "测试甲", "男", "初一", 9, "一分钟跳绳", "400米"])
            workbook.save(path)

            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(len(imported.athletes), 1)
            self.assertEqual({entry.event_name for entry in imported.entries}, {"1分钟跳绳", "400米"})
            self.assertEqual(imported.athletes[0].bib, "10901")
            self.assertEqual(len(imported.report.by_severity(Severity.AUTO_FIXED)), 3)
            project_path = Path(directory) / "报名项目.sportsmeet"
            with ProjectDatabase.create(project_path) as database:
                database.replace_registration(imported.athletes, imported.entries)
                athletes, entries = database.load_registration()
            self.assertEqual(athletes, imported.athletes)
            self.assertEqual(entries, imported.entries)

    def test_teacher_assigned_bibs_are_always_replaced_by_system_numbers(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = []
            for index, name in enumerate(("甲", "乙")):
                path = Path(directory) / f"{index}.xlsx"
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(["号码", "姓名", "性别", "年级", "班级", "项目1"])
                sheet.append([99999, name, "男", "初一", 1, "50米"])
                workbook.save(path)
                paths.append(path)
            imported = import_registration_workbooks(paths)
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(
                {athlete.name: athlete.bib for athlete in imported.athletes},
                {"甲": "10101", "乙": "10103"},
            )
            replaced = [issue for issue in imported.report.issues if issue.code == "IMPORT_014"]
            self.assertEqual(len(replaced), 2)
            self.assertEqual(
                {(issue.original_value, issue.normalized_value) for issue in replaced},
                {("99999", "10101"), ("99999", "10103")},
            )
            self.assertEqual(
                {(issue.source_file_id and imported.report.source_files[issue.source_file_id], issue.source_sheet, issue.source_row)
                 for issue in replaced},
                {("0.xlsx", "Sheet", 2), ("1.xlsx", "Sheet", 2)},
            )

    def test_missing_class_inside_imported_grade_range_is_reported(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "高一报名汇总.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "报名名单"
            sheet.append(list(REGISTRATION_HEADERS))
            for class_number in (*range(1, 8), *range(9, 13)):
                sheet.append([None, f"学生{class_number}", "男", "高一", class_number, "100米", None])
            workbook.save(path)

            imported = import_registration_workbooks([path])
            issue = next(issue for issue in imported.report.issues if issue.code == "IMPORT_013")
            self.assertEqual(issue.severity, Severity.WARNING)
            self.assertIn("高一(8)班", issue.message)
            self.assertEqual(imported.report.source_files[issue.source_file_id], "批量汇总检查")
            self.assertEqual(issue.source_sheet, "跨文件汇总")
            self.assertIsNone(issue.source_row)

    def test_blank_bibs_are_generated_with_legacy_grade_class_and_sex_sequence(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "无号码报名.xlsx"
            workbook = Workbook()
            boys = workbook.active
            boys.title = "男生"
            boys.append(list(REGISTRATION_HEADERS))
            boys.append([None, "男甲", "男", "初一", 1, "50米", None])
            boys.append([None, "男乙", "男", "初一", 1, "400米", None])
            girls = workbook.create_sheet("女生")
            girls.append(list(REGISTRATION_HEADERS))
            girls.append([None, "女甲", "女", "初一", 1, "50米", None])
            girls.append([None, "女乙", "女", "初一", 1, "800米", None])
            workbook.save(path)

            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(
                {athlete.name: athlete.bib for athlete in imported.athletes},
                {"男甲": "10101", "男乙": "10103", "女甲": "10102", "女乙": "10104"},
            )
            self.assertEqual(
                len(imported.report.by_severity(Severity.AUTO_FIXED)), 1
            )

    def test_blank_grade_and_class_are_filled_from_same_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "班主任只填一次年级班级.xlsx"
            workbook = Workbook()
            boys = workbook.active
            boys.title = "男生"
            boys.append(list(REGISTRATION_HEADERS))
            boys.append([None, "男甲", "男", "初一", 6, "50米", None])
            boys.append([None, "男乙", "男", None, None, "400米", None])
            girls = workbook.create_sheet("女生")
            girls.append(list(REGISTRATION_HEADERS))
            girls.append([None, "女甲", "女", None, None, "800米", None])
            workbook.save(path)

            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(
                {athlete.name: (athlete.grade, athlete.class_name, athlete.bib) for athlete in imported.athletes},
                {
                    "男甲": ("初一", "6", "10601"),
                    "男乙": ("初一", "6", "10603"),
                    "女甲": ("初一", "6", "10602"),
                },
            )
            filled = [issue for issue in imported.report.issues if issue.code == "IMPORT_015"]
            self.assertEqual(len(filled), 2)
            self.assertEqual(
                {(issue.source_sheet, issue.source_row) for issue in filled},
                {("\u7537\u751f", 3), ("\u5973\u751f", 2)},
            )
            self.assertTrue(all(issue.severity is Severity.AUTO_FIXED for issue in filled))
            self.assertTrue(all(issue.original_value == {"年级": None, "班级": None} for issue in filled))
            self.assertTrue(all(issue.normalized_value == {"年级": "初一", "班级": "6"} for issue in filled))

    def test_single_grade_and_class_value_can_fill_rows_above_it(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "后面才填年级班级.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(list(REGISTRATION_HEADERS))
            sheet.append([None, "甲", "男", None, None, "50米", None])
            sheet.append([None, "乙", "男", "初二", 5, "400米", None])
            workbook.save(path)

            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(
                {athlete.name: (athlete.grade, athlete.class_name) for athlete in imported.athletes},
                {"甲": ("初二", "5"), "乙": ("初二", "5")},
            )
            filled = [issue for issue in imported.report.issues if issue.code == "IMPORT_015"]
            self.assertEqual([(issue.source_sheet, issue.source_row) for issue in filled], [("Sheet", 2)])

    def test_grade_aliases_are_normalized_and_audited(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "年级多种写法.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(list(REGISTRATION_HEADERS))
            rows = (
                ("甲", "七年级", 1, "初一"),
                ("乙", "7年级", 1, "初一"),
                ("丙", "初一年级", 1, "初一"),
                ("丁", "初中一年级", 1, "初一"),
                ("戊", "八年级", 2, "初二"),
                ("己", "9年级", 3, "初三"),
                ("庚", "高中一年级", 1, "高一"),
                ("辛", "11年级", 2, "高二"),
                ("壬", "高3年级", 3, "高三"),
            )
            for name, source_grade, class_number, _expected in rows:
                sheet.append([None, name, "男", source_grade, class_number, "50米", None])
            workbook.save(path)

            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            expected = {name: expected_grade for name, _source, _class, expected_grade in rows}
            self.assertEqual(
                {athlete.name: athlete.grade for athlete in imported.athletes},
                expected,
            )
            normalized = [issue for issue in imported.report.issues if issue.code == "IMPORT_016"]
            self.assertEqual(len(normalized), len(rows))
            self.assertEqual(
                {(issue.source_row, issue.original_value, issue.normalized_value) for issue in normalized},
                {
                    (row_number, source_grade, expected_grade)
                    for row_number, (_name, source_grade, _class, expected_grade)
                    in enumerate(rows, start=2)
                },
            )

    def test_ambiguous_primary_school_grade_is_not_guessed(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "无法判断的年级.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(list(REGISTRATION_HEADERS))
            sheet.append([None, "甲", "男", "一年级", 1, "50米", None])
            workbook.save(path)

            imported = import_registration_workbooks([path])
            self.assertTrue(imported.report.has_blocking)
            self.assertNotIn("IMPORT_016", {issue.code for issue in imported.report.issues})

    def test_same_name_in_same_class_is_reported_as_warning(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "同名学生报名表.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "男生"
            sheet.append(list(REGISTRATION_HEADERS))
            sheet.append([None, "钟安程", "男", "七年级", 4, "400米", None])
            sheet.append([None, "钟 安程", "男", "初一", 4, "立定跳远", None])
            sheet.append([None, "钟安程", "男", "初一", 5, "50米", None])
            workbook.save(path)

            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            duplicates = [issue for issue in imported.report.issues if issue.code == "IMPORT_017"]
            self.assertEqual(len(duplicates), 1)
            issue = duplicates[0]
            self.assertEqual(issue.severity, Severity.WARNING)
            self.assertEqual((issue.source_sheet, issue.source_row), ("男生", 3))
            self.assertEqual(issue.entity_id, "初一(4)班/钟 安程")
            self.assertEqual(
                issue.normalized_value,
                {
                    "首次出现文件": "同名学生报名表.xlsx",
                    "首次出现工作表": "男生",
                    "首次出现行": 2,
                },
            )

    def test_single_class_outlier_is_corrected_to_dominant_class(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "班级异常报名表.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "女生"
            sheet.append(list(REGISTRATION_HEADERS))
            sheet.append([None, "黄楚翎", "女", "高二", 6, "仰卧起坐", None])
            for index, name in enumerate(("卢屿寻", "黄玉颖", "黄庆颖", "郑秒", "罗晨睿"), start=7):
                sheet.append([None, name, "女", "高二", 11, "800米", None])
            workbook.save(path)

            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual({athlete.class_name for athlete in imported.athletes}, {"11"})
            corrections = [issue for issue in imported.report.issues if issue.code == "IMPORT_018"]
            self.assertEqual(len(corrections), 1)
            self.assertEqual((corrections[0].source_sheet, corrections[0].source_row), ("女生", 2))
            self.assertEqual((corrections[0].original_value, corrections[0].normalized_value), ("6", "11"))
            self.assertNotIn("IMPORT_019", {issue.code for issue in imported.report.issues})

    def test_unresolved_multiple_classes_in_one_workbook_are_reported(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "班级混填.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "女生"
            sheet.append(list(REGISTRATION_HEADERS))
            sheet.append([None, "甲", "女", "高二", 6, "800米", None])
            sheet.append([None, "乙", "女", "高二", 6, "400米", None])
            sheet.append([None, "丙", "女", "高二", 11, "800米", None])
            sheet.append([None, "丁", "女", "高二", 11, "400米", None])
            workbook.save(path)

            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            warnings = [issue for issue in imported.report.issues if issue.code == "IMPORT_019"]
            self.assertEqual(len(warnings), 1)
            issue = warnings[0]
            self.assertEqual(issue.severity, Severity.WARNING)
            self.assertEqual((issue.source_sheet, issue.entity_id), ("跨工作表汇总", "高二"))
            self.assertEqual(issue.normalized_value, {"班级分布": {"6": 2, "11": 2}})
            self.assertEqual(
                [(item["Excel行号"], item["班级"]) for item in issue.original_value],
                [(2, "6"), (3, "6"), (4, "11"), (5, "11")],
            )

    def test_filename_grade_and_class_override_conflicting_cells(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "高二年级11班报名表.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "女生"
            sheet.append(list(REGISTRATION_HEADERS))
            sheet.append([None, "黄楚翎", "女", "高一", 6, "800米", None])
            sheet.append([None, "卢屿寻", "女", "高二", 11, "400米", None])
            workbook.save(path)

            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(
                {athlete.name: (athlete.grade, athlete.class_name) for athlete in imported.athletes},
                {"黄楚翎": ("高二", "11"), "卢屿寻": ("高二", "11")},
            )
            overrides = [issue for issue in imported.report.issues if issue.code == "IMPORT_020"]
            self.assertEqual(len(overrides), 1)
            self.assertEqual((overrides[0].source_sheet, overrides[0].source_row), ("女生", 2))
            self.assertEqual(
                overrides[0].original_value,
                {
                    "年级": {"原值": "高一", "文件名识别值": "高二"},
                    "班级": {"原值": "6", "文件名识别值": "11"},
                },
            )
            self.assertNotIn("IMPORT_019", {issue.code for issue in imported.report.issues})

    def test_grade_wide_filename_keeps_multiple_classes_without_warning(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "高二年级报名表.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(list(REGISTRATION_HEADERS))
            sheet.append([None, "甲", "男", "高一", 1, "100米", None])
            sheet.append([None, "乙", "男", "高二", 2, "200米", None])
            sheet.append([None, "丙", "男", "高二", 3, "400米", None])
            workbook.save(path)

            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual({athlete.grade for athlete in imported.athletes}, {"高二"})
            self.assertEqual({athlete.class_name for athlete in imported.athletes}, {"1", "2", "3"})
            self.assertNotIn("IMPORT_019", {issue.code for issue in imported.report.issues})
            overrides = [issue for issue in imported.report.issues if issue.code == "IMPORT_020"]
            self.assertEqual(len(overrides), 1)
            self.assertEqual(overrides[0].source_row, 2)

    def test_filtered_hidden_rows_are_still_imported_and_checked(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "带筛选的报名表.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "男生"
            sheet.append(list(REGISTRATION_HEADERS))
            sheet.append([None, "张桢越", "男", "初一", 4, "400米", None])
            sheet.append([None, "张桢越", "男", "初一", 4, "立定跳远", None])
            sheet.auto_filter.ref = "A1:G3"
            sheet.row_dimensions[3].hidden = True
            workbook.save(path)

            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(len(imported.athletes), 2)
            duplicates = [issue for issue in imported.report.issues if issue.code == "IMPORT_017"]
            self.assertEqual(len(duplicates), 1)
            self.assertEqual(duplicates[0].source_row, 3)

    def test_more_than_two_distinct_events_is_blocking(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "超额报名.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([*REGISTRATION_HEADERS, "项目3"])
            sheet.append([None, "测试甲", "男", "初一", 1, "50米", "400米", "1000米"])
            workbook.save(path)
            imported = import_registration_workbooks([path])
            self.assertTrue(imported.report.has_blocking)
            self.assertIn("IMPORT_010", {issue.code for issue in imported.report.issues})

    def test_registration_event_shorthands_use_grade_and_sex_configuration(self):
        with tempfile.TemporaryDirectory() as directory:
            junior_path = Path(directory) / "初一7班报名表.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "男生"
            sheet.append(list(REGISTRATION_HEADERS))
            sheet.append([None, "甲", "男", "初一", 7, "跳绳", "跳远"])
            workbook.save(junior_path)

            senior_path = Path(directory) / "高二4班报名表.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "女生"
            sheet.append(list(REGISTRATION_HEADERS))
            sheet.append([None, "乙", "女", "高二", 4, 100, 800])
            workbook.save(senior_path)

            imported = import_registration_workbooks(
                [junior_path, senior_path], ProjectConfig()
            )
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(
                {(entry.group_name, entry.event_name) for entry in imported.entries},
                {
                    ("初一男子组", "1分钟跳绳"),
                    ("初一男子组", "立定跳远"),
                    ("高二女子组", "100米"),
                    ("高二女子组", "800米"),
                },
            )
            normalized = [issue for issue in imported.report.issues if issue.code == "IMPORT_007"]
            self.assertEqual(len(normalized), 4)

    def test_unknown_configured_event_blocks_and_name_without_event_is_excluded(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "初一1班报名表.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(list(REGISTRATION_HEADERS))
            sheet.append([None, "无项目", "男", "初一", 1, None, None])
            sheet.append([None, "错误项目", "男", "初一", 1, "铅球", None])
            workbook.save(path)

            imported = import_registration_workbooks([path], ProjectConfig())
            self.assertTrue(imported.report.has_blocking)
            self.assertEqual(imported.athletes, [])
            self.assertIn("IMPORT_021", {issue.code for issue in imported.report.issues})
            empty_issues = [issue for issue in imported.report.issues if issue.code == "IMPORT_022"]
            self.assertEqual(len(empty_issues), 2)

    def test_registration_source_snapshot_detects_modified_and_added_files(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "报名表"
            root.mkdir()
            source = root / "初一1班报名表.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(list(REGISTRATION_HEADERS))
            sheet.append([None, "甲", "男", "初一", 1, "50米", None])
            workbook.save(source)

            imported = import_registration_workbooks([source], ProjectConfig())
            project_path = Path(directory) / "项目.sportsmeet"
            with ProjectDatabase.create(project_path) as database:
                database.replace_registration(
                    imported.athletes,
                    imported.entries,
                    imported.source_snapshots,
                    root,
                )
                self.assertEqual(database.registration_source_changes(), [])

                workbook = load_workbook(source)
                workbook.active["F2"] = "400米"
                workbook.save(source)
                added = root / "初一2班报名表.xlsx"
                workbook = Workbook()
                workbook.active.append(list(REGISTRATION_HEADERS))
                workbook.save(added)

                changes = database.registration_source_changes()
                self.assertTrue(any(item.startswith("已修改：") for item in changes))
                self.assertTrue(any(item.startswith("新增文件：") for item in changes))

    def test_available_lane_parser_supports_ranges_and_gaps(self):
        self.assertEqual(parse_available_lanes("2-4,6,8"), (2, 3, 4, 6, 8))

    def test_imports_structured_schedule(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "日程.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["日期", "时间段", "比赛类型", "项目名称", "组别", "项目", "人数", "组数", "比赛时间", "可用道次"])
            sheet.append(["10月23日", "上午", "径赛", "初一男子组50米预赛", "初一男子组", "50米", "27人", "4组", "9:30", "2-8"])
            workbook.save(path)
            imported = import_schedule_workbook(path)
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(len(imported.rounds), 1)
            self.assertEqual(imported.rounds[0].round_type, RoundType.PRELIMINARY)
            self.assertEqual(imported.rounds[0].available_lanes, tuple(range(2, 9)))

    def test_imports_legacy_team_rows_and_suggests_preliminary_heats(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "旧日程.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["竞赛项目", "人数", "类型", "时间", "上下午", "日期", "序号", "组别", "项目"])
            sheet.append(["初一男子组50米预赛", 27, "径赛", "9:30", "上午", "10-11", 1, "初一男子组", "50米"])
            sheet.append(["初一年级女子4 X 100米接力赛", "14队", "径赛", "13:30", "下午", "10-12", 2, "初一年级组", None])
            workbook.save(path)

            imported = import_schedule_workbook(path)
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(len(imported.rounds), 2)
            self.assertEqual(imported.rounds[0].heat_count, 4)
            self.assertEqual(imported.rounds[1].group_name, "初一女子组")
            self.assertEqual(imported.rounds[1].event_name, "4×100米接力")
            self.assertEqual(imported.rounds[1].event_type, EventType.TEAM)

    def test_schedule_infers_user_friendly_default_score_units(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "单位日程.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["项目名称", "组别", "项目", "比赛类型", "人数", "组数", "比赛时间"])
            sheet.append(["初一男子组50米预赛", "初一男子组", "50米", "径赛", 8, 1, "9:00"])
            sheet.append(["初一男子组800米预赛", "初一男子组", "800米", "径赛", 8, 1, "9:20"])
            sheet.append(["初一女子组跳远预赛", "初一女子组", "跳远", "田赛", 8, 1, "9:40"])
            sheet.append(["初一男子组引体向上预赛", "初一男子组", "引体向上", "田赛", 8, 1, "10:00"])
            workbook.save(path)

            imported = import_schedule_workbook(path)
            units = {item.event_name: item.default_input_unit for item in imported.rounds}
            self.assertEqual(units["50米"], "SECOND")
            self.assertEqual(units["800米"], "MINUTE")
            self.assertEqual(units["跳远"], "METER")
            self.assertEqual(units["引体向上"], "COUNT")


class DatabaseTests(unittest.TestCase):
    def test_project_round_trip_and_backup_retention(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.sportsmeet"
            with ProjectDatabase.create(path, ProjectConfig(backup_retention=2)) as database:
                self.assertEqual(database.load_config().lane_priority, (4, 5, 3, 6, 2, 7, 1, 8))
                database.backup()
                database.backup()
                database.backup()
            self.assertEqual(len(list((Path(directory) / "test_backups").glob("*.sportsmeet"))), 2)

    def test_schedule_round_trip(self):
        with tempfile.TemporaryDirectory() as directory:
            schedule_path = Path(directory) / "日程.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["项目名称", "组别", "项目", "比赛类型", "人数", "组数", "比赛时间"])
            sheet.append(["初一男子组50米预赛", "初一男子组", "50米", "径赛", 16, 2, "9:30"])
            workbook.save(schedule_path)
            imported = import_schedule_workbook(schedule_path)

            project_path = Path(directory) / "test.sportsmeet"
            with ProjectDatabase.create(project_path) as database:
                database.replace_schedule(imported.rounds)
                loaded = database.load_schedule()
                assignment = HeatAssignment(
                    "assignment-1", imported.rounds[0].id, "athlete-1", 1, 4,
                    random_seed=2026,
                )
                database.replace_heat_assignments([assignment], 2026)
                loaded_assignments = database.load_heat_assignments()
            self.assertEqual(loaded, imported.rounds)
            self.assertEqual(loaded_assignments, [assignment])

    def test_results_qualifications_and_final_lanes_round_trip(self):
        with tempfile.TemporaryDirectory() as directory:
            registration = import_registration_workbooks([
                create_grouping_registration_sample(Path(directory) / "报名.xlsx")
            ])
            schedule = import_schedule_workbook(
                create_grouping_schedule_sample(Path(directory) / "日程.xlsx")
            )
            preliminary = next(
                item for item in schedule.rounds if item.round_type is RoundType.PRELIMINARY
            )
            final_round = next(
                item for item in schedule.rounds if item.round_type is RoundType.FINAL
            )
            generated = generate_heat_assignments(
                ProjectConfig(), registration.athletes, registration.entries,
                schedule.rounds, seed=2026,
            )
            preliminary_assignments = sorted(
                (
                    item for item in generated.assignments
                    if item.event_round_id == preliminary.id
                ),
                key=lambda item: (item.heat_no, item.lane or 0),
            )
            results = [
                Result(
                    item.participant_id, item.heat_no, f"{7 + index / 100:.2f}",
                    PerformanceKind.TIME, 7000 + index * 10, "MILLISECOND",
                    f"{7 + index / 100:.2f}",
                )
                for index, item in enumerate(preliminary_assignments)
            ]
            computed = compute_track_final(
                preliminary, final_round, results, ProjectConfig()
            )
            project_path = Path(directory) / "成绩项目.sportsmeet"
            with ProjectDatabase.create(project_path) as database:
                database.replace_registration(registration.athletes, registration.entries)
                database.replace_schedule(schedule.rounds)
                database.replace_heat_assignments(generated.assignments, 2026)
                database.save_track_final(
                    preliminary.id, final_round.id, results,
                    computed.qualifications, computed.assignments,
                )
                self.assertEqual(database.load_results(preliminary.id), results)
                self.assertEqual(len(database.load_qualifications(preliminary.id)), 8)
                final_assignments = [
                    item for item in database.load_heat_assignments()
                    if item.event_round_id == final_round.id
                ]
                self.assertEqual(len(final_assignments), 8)
                best_id = min(results, key=lambda item: item.canonical_value).participant_id
                self.assertEqual(
                    next(item.lane for item in final_assignments if item.participant_id == best_id),
                    4,
                )


class WorkbookTemplateTests(unittest.TestCase):
    def test_school_level_templates_keep_sex_specific_event_lists(self):
        with tempfile.TemporaryDirectory() as directory:
            path = create_school_level_registration_template(
                Path(directory) / "初中报名模板.xlsx",
                "初中",
                ("50米", "男生测试项目"),
                ("800米", "女生测试项目"),
            )
            workbook = load_workbook(path, read_only=False, data_only=False)
            self.assertEqual(workbook.active.title, "男生")
            self.assertEqual(workbook["_下拉数据"].sheet_state, "hidden")
            self.assertEqual(
                workbook["男生"]["A1"].value,
                "示例学校田径运动会报名表（初中年级用）",
            )
            self.assertIn("每班每项不得少于一人，可兼接力", workbook["男生"]["A2"].value)
            self.assertEqual(tuple(cell.value for cell in workbook["男生"][5]), REGISTRATION_HEADERS)
            self.assertEqual(workbook["男生"]["A4"].value, "男生项目：50米、男生测试项目")
            self.assertEqual(workbook["女生"]["A4"].value, "女生项目：800米、女生测试项目")
            male_formulas = [item.formula1 for item in workbook["男生"].data_validations.dataValidation]
            female_formulas = [item.formula1 for item in workbook["女生"].data_validations.dataValidation]
            self.assertIn("'_下拉数据'!$A$2:$A$3", male_formulas)
            self.assertIn("'_下拉数据'!$B$2:$B$3", female_formulas)
            workbook.close()
            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(imported.report.issues, [])

    def test_registration_template_has_dropdowns_and_importable_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            path = create_registration_template(Path(directory) / "报名模板.xlsx")
            workbook = load_workbook(path, read_only=False, data_only=True)
            sheet = workbook["报名信息"]
            self.assertEqual(tuple(cell.value for cell in sheet[1]), REGISTRATION_HEADERS)
            self.assertEqual(len(sheet.data_validations.dataValidation), 3)
            self.assertIn("填写说明（不要导入）", workbook.sheetnames)
            workbook.close()
            imported = import_registration_workbooks([path])
            self.assertEqual(imported.athletes, [])
            self.assertFalse(imported.report.has_blocking)

    def test_level_registration_template_separates_sexes_and_uses_configured_events(self):
        with tempfile.TemporaryDirectory() as directory:
            path = create_school_level_registration_template(
                Path(directory) / "初中报名模板.xlsx",
                "初中",
                ("50米", "引体向上"),
                ("50米", "仰卧起坐"),
            )
            workbook = load_workbook(path, read_only=False, data_only=False)
            self.assertEqual(workbook.sheetnames, ["_下拉数据", "男生", "女生"])
            self.assertEqual(workbook["_下拉数据"].sheet_state, "hidden")
            self.assertEqual(tuple(cell.value for cell in workbook["男生"][5]), REGISTRATION_HEADERS)
            self.assertEqual(workbook["男生"]["A4"].value, "男生项目：50米、引体向上")
            self.assertEqual(workbook["女生"]["A4"].value, "女生项目：50米、仰卧起坐")
            boy_validations = workbook["男生"].data_validations.dataValidation
            girl_validations = workbook["女生"].data_validations.dataValidation
            self.assertEqual(boy_validations[-1].formula1, "'_下拉数据'!$A$2:$A$3")
            self.assertEqual(girl_validations[-1].formula1, "'_下拉数据'!$B$2:$B$3")
            workbook.close()
            imported = import_registration_workbooks([path])
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(imported.report.issues, [])

    def test_registration_event_settings_round_trip_with_project(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "项目.sportsmeet"
            config = ProjectConfig(
                junior_boys_events=("50米", "跳高"),
                junior_boys_event_formats={"50米": ROUND_MODE_PRELIMINARY_FINAL},
                meeting_start_date="2026-10-03",
                meeting_end_date="2026-10-05",
                booklet_title="测试运动会秩序册",
                competition_rules_text="测试规程正文",
            )
            with ProjectDatabase.create(path, config) as database:
                loaded = database.load_config()
                self.assertEqual(loaded.junior_boys_events, ("50米", "跳高"))
                self.assertEqual(
                    loaded.junior_boys_event_formats,
                    {"50米": ROUND_MODE_PRELIMINARY_FINAL},
                )
                self.assertEqual(loaded.meeting_start_date, "2026-10-03")
                self.assertEqual(loaded.meeting_end_date, "2026-10-05")
                self.assertEqual(loaded.senior_girls_events, ProjectConfig().senior_girls_events)
                self.assertEqual(loaded.booklet_title, "测试运动会秩序册")
                self.assertEqual(loaded.competition_rules_text, "测试规程正文")

    def test_schedule_template_has_blank_input_sheet_and_separate_examples(self):
        with tempfile.TemporaryDirectory() as directory:
            path = create_schedule_template(Path(directory) / "日程模板.xlsx")
            workbook = load_workbook(path, read_only=False, data_only=True)
            self.assertEqual(workbook.sheetnames, ["比赛日程", "填写说明", "填写示例（不要导入）"])
            schedule = workbook["比赛日程"]
            self.assertEqual(tuple(cell.value for cell in schedule[1]), SCHEDULE_HEADERS)
            self.assertEqual(schedule.max_row, 1)
            self.assertGreater(workbook["填写示例（不要导入）"].max_row, 1)
            self.assertEqual(len(schedule.data_validations.dataValidation), 2)
            workbook.close()
            imported = import_schedule_workbook(path)
            self.assertEqual(imported.rounds, [])
            self.assertFalse(imported.report.has_blocking)

    def test_schedule_test_sample_has_predictable_import_result(self):
        with tempfile.TemporaryDirectory() as directory:
            path = create_schedule_test_sample(Path(directory) / "日程测试样例.xlsx")
            imported = import_schedule_workbook(path)
            self.assertEqual(len(imported.rounds), 3)
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(imported.rounds[0].heat_count, 3)
            self.assertEqual(len(imported.report.by_severity(Severity.AUTO_FIXED)), 1)

    def test_grouping_samples_pass_validation_and_generate_expected_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            registration = import_registration_workbooks([
                create_grouping_registration_sample(Path(directory) / "报名.xlsx")
            ])
            schedule = import_schedule_workbook(
                create_grouping_schedule_sample(Path(directory) / "日程.xlsx")
            )
            report = validate_project(
                ProjectConfig(), registration.athletes, registration.entries, schedule.rounds
            )
            self.assertFalse(report.has_blocking)
            generated = generate_heat_assignments(
                ProjectConfig(), registration.athletes, registration.entries, schedule.rounds, seed=2026
            )
            self.assertFalse(generated.report.has_blocking)
            self.assertEqual(len(generated.assignments), 25)


if __name__ == "__main__":
    unittest.main()
