import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from sport_meeting.core.importers import import_schedule_workbook
from sport_meeting.core.models import (
    ROUND_MODE_PRELIMINARY_FINAL,
    ROUND_MODE_TIMED_FINAL,
    Athlete,
    Entry,
    EventType,
    ProjectConfig,
    TEAM_COUNT_ONE_PER_CLASS,
    TEAM_GROUP_MIXED,
    TEAM_ROUND_MODE_KNOCKOUT,
    TeamEventConfig,
)
from sport_meeting.core.schedule_planner import (
    ScheduleCapacityError,
    build_standard_schedule,
    guess_event_round_mode,
)
from sport_meeting.core.workbook_templates import create_schedule_template


class SchedulePlannerTests(unittest.TestCase):
    def _config(self, **changes) -> ProjectConfig:
        values = {
            "meeting_start_date": "2026-10-23",
            "meeting_end_date": "2026-10-24",
            "junior_boys_events": ("50米",),
            "junior_girls_events": ("50米",),
            "senior_boys_events": ("100米", "200米"),
            "senior_girls_events": ("100米", "200米"),
        }
        values.update(changes)
        return ProjectConfig(**values)

    def test_expands_only_grades_present_in_registration_and_uses_opening_start(self):
        athletes = [
            Athlete("j1", "10101", "初一学生", "男", "初一", "1"),
            Athlete("j2", "20101", "初二学生", "男", "初二", "1"),
            Athlete("s1", "40101", "高一学生", "女", "高一", "1"),
            Athlete("s2", "50101", "高二学生", "女", "高二", "1"),
        ]
        rows = build_standard_schedule(self._config(), athletes)
        groups = {row.group_name for row in rows}
        self.assertTrue({"初一男子组", "初二男子组", "高一女子组", "高二女子组"} <= groups)
        self.assertFalse(any(group.startswith("初三") for group in groups))
        self.assertFalse(any(group.startswith("高三") for group in groups))
        first_track = next(row for row in rows if row.event_type.value == "TRACK")
        self.assertEqual(first_track.competition_date.isoformat(), "2026-10-23")
        self.assertEqual(first_track.start_time.strftime("%H:%M"), "09:30")
        self.assertTrue(any("预赛" in row.event_title for row in rows))
        self.assertTrue(any("决赛" in row.event_title for row in rows))

    def test_registration_counts_drive_heat_count_and_template_values(self):
        athletes = [
            Athlete(f"a{index}", f"101{index:02d}", f"学生{index}", "男", "初一", "1")
            for index in range(1, 18)
        ]
        entries = [
            Entry(f"e{index}", athlete.id, "50米", "初一男子组")
            for index, athlete in enumerate(athletes, start=1)
        ]
        rows = build_standard_schedule(self._config(), athletes, entries)
        preliminary = next(
            row for row in rows
            if row.group_name == "初一男子组" and row.event_name == "50米" and "预赛" in row.event_title
        )
        final = next(
            row for row in rows
            if row.group_name == "初一男子组" and row.event_name == "50米" and row.event_title.endswith("决赛")
        )
        self.assertEqual((preliminary.declared_count, preliminary.heat_count), (17, 3))
        self.assertEqual((final.declared_count, final.heat_count), (8, 1))

        with tempfile.TemporaryDirectory() as directory:
            path = create_schedule_template(Path(directory) / "自动日程.xlsx", rows)
            workbook = load_workbook(path, data_only=True)
            schedule = workbook["比赛日程"]
            values = list(schedule.iter_rows(min_row=2, values_only=True))
            target = next(row for row in values if row[4] == preliminary.event_title)
            self.assertEqual(target[7:9], (17, 3))
            self.assertEqual(target[9].strftime("%H:%M"), preliminary.start_time.strftime("%H:%M"))
            workbook.close()
            imported = import_schedule_workbook(path)
            self.assertFalse(imported.report.has_blocking)
            self.assertEqual(len(imported.rounds), len(rows))

    def test_round_mode_is_guessed_but_saved_choice_takes_priority(self):
        self.assertEqual(guess_event_round_mode("50米"), ROUND_MODE_PRELIMINARY_FINAL)
        self.assertEqual(guess_event_round_mode("400米"), ROUND_MODE_TIMED_FINAL)
        athlete = Athlete("a1", "10101", "学生", "男", "初一", "1")

        timed_final_rows = build_standard_schedule(
            self._config(
                junior_boys_event_formats={"50米": ROUND_MODE_TIMED_FINAL},
            ),
            [athlete],
        )
        target_titles = [
            row.event_title for row in timed_final_rows
            if row.group_name == "初一男子组" and row.event_name == "50米"
        ]
        self.assertEqual(len(target_titles), 1)
        self.assertTrue(target_titles[0].endswith("预决赛"))

        preliminary_rows = build_standard_schedule(
            self._config(
                junior_boys_events=("400米",),
                junior_boys_event_formats={"400米": ROUND_MODE_PRELIMINARY_FINAL},
            ),
            [athlete],
        )
        target_titles = [
            row.event_title for row in preliminary_rows
            if row.group_name == "初一男子组" and row.event_name == "400米"
        ]
        self.assertEqual(len(target_titles), 2)
        self.assertTrue(any(title.endswith("预赛") for title in target_titles))
        self.assertTrue(any(title.endswith("决赛") for title in target_titles))

    def test_timed_final_sprints_split_by_lanes_but_distance_runs_stay_mass_start(self):
        athletes = [
            Athlete(f"a{index}", f"501{index:02d}", f"学生{index}", "男", "高二", "1")
            for index in range(1, 25)
        ]
        entries = [
            Entry(f"s{index}", athlete.id, "100米", "高二男子组")
            for index, athlete in enumerate(athletes, start=1)
        ] + [
            Entry(f"d{index}", athlete.id, "1500米", "高二男子组")
            for index, athlete in enumerate(athletes, start=1)
        ]
        config = self._config(
            senior_boys_events=("100米", "1500米"),
            senior_boys_event_formats={
                "100米": ROUND_MODE_TIMED_FINAL,
                "1500米": ROUND_MODE_TIMED_FINAL,
            },
        )
        rows = build_standard_schedule(config, athletes, entries)
        sprint = next(
            row for row in rows
            if row.group_name == "高二男子组" and row.event_name == "100米"
        )
        distance = next(
            row for row in rows
            if row.group_name == "高二男子组" and row.event_name == "1500米"
        )
        self.assertEqual((sprint.declared_count, sprint.heat_count), (24, 3))
        self.assertEqual((distance.declared_count, distance.heat_count), (24, 1))

    def test_scheduler_allows_track_and_field_to_run_in_parallel(self):
        athlete = Athlete("a1", "10101", "兼项学生", "男", "初一", "1")
        entries = [
            Entry("e1", athlete.id, "50米", "初一男子组"),
            Entry("e2", athlete.id, "立定跳远", "初一男子组"),
        ]
        rows = build_standard_schedule(
            self._config(
                junior_boys_events=("50米", "立定跳远"),
                junior_boys_event_formats={
                    "50米": ROUND_MODE_TIMED_FINAL,
                    "立定跳远": ROUND_MODE_TIMED_FINAL,
                },
            ),
            [athlete],
            entries,
        )
        track = next(
            row for row in rows
            if row.group_name == "初一男子组" and row.event_name == "50米"
        )
        field = next(
            row for row in rows
            if row.group_name == "初一男子组" and row.event_name == "立定跳远"
        )
        self.assertEqual(field.start_time.strftime("%H:%M"), "09:30")
        self.assertLess(track.start_time.strftime("%H:%M"), "10:00")

    def test_balances_second_day_with_200m_and_later_parallel_tug_of_war(self):
        athletes = [
            Athlete("j1", "10101", "初一学生", "男", "初一", "1"),
            Athlete("j2", "20101", "初二学生", "男", "初二", "1"),
            Athlete("s1", "40101", "高一学生", "男", "高一", "1"),
            Athlete("s2", "50101", "高二学生", "女", "高二", "1"),
        ]
        entries = [
            Entry("e1", "s1", "200米", "高一男子组"),
            Entry("e2", "s2", "200米", "高二女子组"),
            Entry("e3", "s1", "1500米", "高一男子组"),
            Entry("e4", "s2", "800米", "高二女子组"),
        ]
        rows = build_standard_schedule(
            self._config(
                junior_boys_events=("50米",),
                junior_girls_events=("50米",),
                senior_boys_events=("200米", "1500米"),
                senior_girls_events=("200米", "800米"),
                senior_boys_event_formats={"200米": ROUND_MODE_PRELIMINARY_FINAL},
                senior_girls_event_formats={"200米": ROUND_MODE_PRELIMINARY_FINAL},
                team_events=(
                    TeamEventConfig(
                        "拔河比赛",
                        TEAM_GROUP_MIXED,
                        TEAM_COUNT_ONE_PER_CLASS,
                        TEAM_ROUND_MODE_KNOCKOUT,
                    ),
                ),
            ),
            athletes,
            entries,
        )

        preliminaries = [
            row for row in rows
            if row.event_name == "200米" and row.event_title.endswith("预赛")
        ]
        finals = [
            row for row in rows
            if row.event_name == "200米" and row.event_title.endswith("决赛")
        ]
        distance_runs = [row for row in rows if row.event_name in {"800米", "1500米"}]
        self.assertTrue(preliminaries)
        self.assertTrue(finals)
        self.assertTrue(all(row.competition_date.isoformat() == "2026-10-24" for row in preliminaries))
        self.assertTrue(all(row.period == "上午" for row in preliminaries))
        self.assertGreater(min(row.start_time for row in preliminaries), max(row.start_time for row in distance_runs))
        self.assertTrue(all(row.competition_date.isoformat() == "2026-10-24" for row in finals))
        self.assertTrue(all(row.period == "下午" for row in finals))

        tug_rows = [row for row in rows if row.event_name == "拔河比赛"]
        self.assertTrue(all(row.declared_count == 2 for row in tug_rows))
        self.assertTrue(all(row.heat_count == 1 for row in tug_rows))
        self.assertTrue(all("冠亚军赛" in row.event_title for row in tug_rows))
        self.assertEqual(
            sorted({row.start_time.strftime("%H:%M") for row in tug_rows}),
            ["09:30", "10:00"],
        )
        tug_by_group = {row.group_name: row.start_time.strftime("%H:%M") for row in tug_rows}
        self.assertEqual(tug_by_group["初一年级组"], tug_by_group["高一年级组"])
        self.assertEqual(tug_by_group["初二年级组"], tug_by_group["高二年级组"])

    def test_track_finals_respect_recovery_time(self):
        athlete = Athlete("a1", "10101", "学生", "男", "初一", "1")
        entries = [
            Entry("e1", athlete.id, "50米", "初一男子组"),
            Entry("e2", athlete.id, "400米", "初一男子组"),
        ]
        rows = build_standard_schedule(
            self._config(
                junior_boys_events=("50米", "400米"),
                junior_girls_events=("50米",),
                senior_boys_events=("100米",),
                senior_girls_events=("100米",),
                junior_boys_event_formats={
                    "50米": ROUND_MODE_PRELIMINARY_FINAL,
                    "400米": ROUND_MODE_PRELIMINARY_FINAL,
                },
                team_events=(),
            ),
            [athlete],
            entries,
        )
        by_event = {}
        for row in rows:
            by_event.setdefault(row.event_name, {})[
                "preliminary" if row.event_title.endswith("预赛") else "final"
            ] = row

        sprint_preliminary = by_event["50米"]["preliminary"]
        sprint_final = by_event["50米"]["final"]
        sprint_start_gap = (
            sprint_final.start_time.hour * 60
            + sprint_final.start_time.minute
            - sprint_preliminary.start_time.hour * 60
            - sprint_preliminary.start_time.minute
        )
        self.assertGreaterEqual(sprint_start_gap, 55)  # 10 分钟比赛 + 45 分钟休息

        lap_preliminary = by_event["400米"]["preliminary"]
        lap_final = by_event["400米"]["final"]
        self.assertEqual(lap_preliminary.competition_date, lap_final.competition_date)
        self.assertEqual(lap_preliminary.period, "上午")
        self.assertEqual(lap_final.period, "下午")

    def test_400m_is_split_across_days_without_splitting_each_preliminary_final_pair(self):
        athletes = [
            Athlete("j1", "10101", "初一学生", "男", "初一", "1"),
            Athlete("j2", "20101", "初二学生", "男", "初二", "1"),
            Athlete("s1", "40101", "高一学生", "男", "高一", "1"),
            Athlete("s2", "50101", "高二学生", "男", "高二", "1"),
        ]
        preliminary_final_400 = {"400米": ROUND_MODE_PRELIMINARY_FINAL}
        rows = build_standard_schedule(
            ProjectConfig(
                meeting_start_date="2026-10-23",
                meeting_end_date="2026-10-24",
                junior_boys_event_formats=preliminary_final_400,
                junior_girls_event_formats=preliminary_final_400,
                senior_boys_event_formats=preliminary_final_400,
                senior_girls_event_formats=preliminary_final_400,
            ),
            athletes,
        )
        by_group = {}
        for row in rows:
            if row.event_name != "400米":
                continue
            by_group.setdefault(row.group_name, {})[
                "preliminary" if row.event_title.endswith("预赛") else "final"
            ] = row

        preliminary_dates = {
            pair["preliminary"].competition_date for pair in by_group.values()
        }
        self.assertEqual(len(preliminary_dates), 2)
        for pair in by_group.values():
            self.assertEqual(pair["preliminary"].period, "上午")
            self.assertEqual(pair["final"].period, "下午")
            self.assertEqual(
                pair["preliminary"].competition_date,
                pair["final"].competition_date,
            )
        latest_start_by_date = {
            competition_date: max(
                row.start_time for row in rows if row.competition_date == competition_date
            )
            for competition_date in {row.competition_date for row in rows}
        }
        latest_minutes = [value.hour * 60 + value.minute for value in latest_start_by_date.values()]
        self.assertLessEqual(max(latest_minutes) - min(latest_minutes), 60)

    def test_registration_without_entries_leaves_people_and_heats_blank(self):
        athletes = [Athlete("a1", "10101", "学生", "男", "初一", "1")]
        rows = build_standard_schedule(self._config(), athletes)
        individual_rows = [row for row in rows if row.event_type is not EventType.TEAM]
        self.assertTrue(all(row.declared_count is None for row in individual_rows))
        self.assertTrue(all(row.heat_count is None for row in individual_rows))
        self.assertTrue(
            all(row.group_name.startswith("初一") for row in rows if row.group_name != "教工组")
        )

    def test_collective_events_use_actual_classes_and_keep_manual_counts_blank(self):
        athletes = [
            Athlete("a1", "10101", "甲", "男", "初一", "1"),
            Athlete("a2", "10202", "乙", "女", "初一", "2"),
        ]
        rows = build_standard_schedule(self._config(), athletes)

        boys_relay = next(
            row for row in rows
            if row.event_name == "4×100米接力" and row.group_name == "初一男子组"
        )
        girls_relay = next(
            row for row in rows
            if row.event_name == "4×100米接力" and row.group_name == "初一女子组"
        )
        collective_rope = next(
            row for row in rows
            if row.event_name == "1分钟集体跳绳" and row.group_name == "初一年级组"
        )
        tug_of_war = next(
            row for row in rows
            if row.event_name == "拔河比赛" and row.group_name == "初一年级组"
        )
        staff = next(row for row in rows if row.group_name == "教工组")

        self.assertEqual((boys_relay.declared_count, boys_relay.heat_count), (2, 1))
        self.assertEqual((girls_relay.declared_count, girls_relay.heat_count), (2, 1))
        self.assertEqual((collective_rope.declared_count, collective_rope.heat_count), (2, 2))
        self.assertIsNone(tug_of_war.declared_count)
        self.assertIn("淘汰赛", tug_of_war.event_title)
        self.assertIsNone(staff.declared_count)

        with tempfile.TemporaryDirectory() as directory:
            path = create_schedule_template(Path(directory) / "含集体项目的标准日程.xlsx", rows)
            workbook = load_workbook(path, data_only=True)
            exported_rows = list(
                workbook["比赛日程"].iter_rows(min_row=2, values_only=True)
            )
            self.assertTrue(
                any(row[2] == "集体项目" and row[6] == "4×100米接力" for row in exported_rows)
            )
            self.assertTrue(
                any(row[2] == "集体项目" and row[6] == "拔河比赛" for row in exported_rows)
            )
            workbook.close()

    def test_without_registration_requires_import_first(self):
        with self.assertRaisesRegex(ValueError, "请先批量导入报名表"):
            build_standard_schedule(self._config())

    def test_later_competition_day_starts_at_eight(self):
        athletes = [
            Athlete("a1", "10101", "初一学生", "男", "初一", "1"),
            Athlete("a2", "20101", "初二学生", "男", "初二", "1"),
            Athlete("a3", "40101", "高一学生", "男", "高一", "1"),
            Athlete("a4", "50101", "高二学生", "男", "高二", "1"),
        ]
        rows = build_standard_schedule(
            ProjectConfig(
                meeting_start_date="2026-10-23",
                meeting_end_date="2026-10-24",
            ),
            athletes,
        )
        second_day_rows = [
            row for row in rows
            if row.competition_date.isoformat() == "2026-10-24" and row.period == "上午"
        ]
        self.assertTrue(second_day_rows)
        self.assertEqual(min(row.start_time for row in second_day_rows).strftime("%H:%M"), "08:00")

    def test_reports_missing_slots_when_selected_dates_are_too_short(self):
        many_events = tuple(f"50米项目{index}" for index in range(1, 31))
        config = self._config(
            meeting_end_date="2026-10-23",
            junior_boys_events=many_events,
            junior_girls_events=many_events,
            senior_boys_events=many_events,
            senior_girls_events=many_events,
        )
        with self.assertRaises(ScheduleCapacityError) as caught:
            build_standard_schedule(
                config,
                [Athlete("a1", "10101", "初一学生", "男", "初一", "1")],
            )
        self.assertGreater(caught.exception.missing_slots, 0)
        self.assertIn("还差", str(caught.exception))


if __name__ == "__main__":
    unittest.main()
