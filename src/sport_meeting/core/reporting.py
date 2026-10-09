from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Participant, Qualification, Result, Severity, ValidationReport
from .word_layout import fit_single_line_font, normalize_single_line_text, prevent_cell_text_wrap


@dataclass(frozen=True, slots=True)
class FinalMaterial:
    event_title: str
    scheduled_time: str
    heats: list[list[tuple[str, int]]]
    position_header: str = "道次"


@dataclass(frozen=True, slots=True)
class FinalRankingMaterial:
    event_title: str
    ranked_results: list[tuple[int | None, Participant, Result]]


@dataclass(frozen=True, slots=True)
class OnSitePackageFiles:
    assignment_sheet: Path
    check_in_sheet: Path
    official_sheet: Path
    broadcast_script: Path


def _set_font(run, name: str = "微软雅黑", size: int = 10, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def _configure_a4(document: Document) -> None:
    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)


def _heading(document: Document, text: str, size: int = 16) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(paragraph.add_run(text), size=size, bold=True)


def _set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _add_grid_table(document: Document, headers: list[str], rows: Iterable[list[str]]):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    _set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(cell.paragraphs[0].add_run(value), bold=True)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            text = normalize_single_line_text(str(value))
            fit = fit_single_line_font(text, 6.0)
            _set_font(cells[index].paragraphs[0].add_run(text), size=fit.font_size)
            if fit.fits:
                prevent_cell_text_wrap(cells[index])
    return table


def render_final_assignments(
    event_title: str,
    scheduled_time: str,
    heats: list[list[tuple[str, int]]],
    participants: dict[str, Participant],
    output_path: str | Path,
    position_header: str = "道次",
) -> Path:
    document = Document()
    _configure_a4(document)
    _append_final_material(
        document,
        FinalMaterial(event_title, scheduled_time, heats, position_header),
        participants,
    )
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


def render_final_assignments_batch(
    materials: list[FinalMaterial],
    participants: dict[str, Participant],
    output_path: str | Path,
) -> Path:
    if not materials:
        raise ValueError("没有可生成的决赛名单")
    document = Document()
    _configure_a4(document)
    for index, material in enumerate(materials):
        if index:
            document.add_page_break()
        _append_final_material(document, material, participants)
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


def render_on_site_package(
    materials: list[FinalMaterial],
    participants: dict[str, Participant],
    output_directory: str | Path,
) -> OnSitePackageFiles:
    if not materials:
        raise ValueError("没有可生成的决赛现场材料")
    target = Path(output_directory)
    target.mkdir(parents=True, exist_ok=True)
    files = OnSitePackageFiles(
        assignment_sheet=target / "01-决赛分组表.docx",
        check_in_sheet=target / "02-检录单.docx",
        official_sheet=target / "03-发令员与裁判员用表.docx",
        broadcast_script=target / "04-广播稿.txt",
    )
    render_final_assignments_batch(materials, participants, files.assignment_sheet)
    _render_check_in_sheets(materials, participants, files.check_in_sheet)
    _render_official_sheets(materials, participants, files.official_sheet)
    _render_broadcast_script(materials, participants, files.broadcast_script)
    return files


def _render_check_in_sheets(
    materials: list[FinalMaterial], participants: dict[str, Participant], output_path: Path
) -> None:
    document = Document()
    _configure_a4(document)
    for material_index, material in enumerate(materials):
        if material_index:
            document.add_page_break()
        _heading(document, f"{material.event_title} 决赛检录单")
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(subtitle.add_run(f"比赛时间：{material.scheduled_time}"), size=11, bold=True)
        for heat_index, assignments in enumerate(material.heats, start=1):
            _heading(document, f"第{heat_index}组", size=12)
            rows = [
                [
                    str(position), participant.bib, participant.name, participant.unit,
                    "□", "", "",
                ]
                for participant_id, position in sorted(assignments, key=lambda item: item[1])
                for participant in [participants[participant_id]]
            ]
            _add_grid_table(
                document,
                [material.position_header, "号码", "姓名", "单位", "到检", "检录时间", "备注"],
                rows,
            )
    document.save(output_path)


def _render_official_sheets(
    materials: list[FinalMaterial], participants: dict[str, Participant], output_path: Path
) -> None:
    document = Document()
    _configure_a4(document)
    for material_index, material in enumerate(materials):
        if material_index:
            document.add_page_break()
        sheet_name = "发令员道次表" if material.position_header == "道次" else "裁判员出场顺序表"
        _heading(document, f"{material.event_title} {sheet_name}", size=18)
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(subtitle.add_run(f"比赛时间：{material.scheduled_time}"), size=12, bold=True)
        for heat_index, assignments in enumerate(material.heats, start=1):
            if len(material.heats) > 1:
                _heading(document, f"第{heat_index}组", size=13)
            rows = [
                [str(position), participants[participant_id].bib, participants[participant_id].name]
                for participant_id, position in sorted(assignments, key=lambda item: item[1])
            ]
            table = _add_grid_table(document, [material.position_header, "号码", "姓名"], rows)
            for row in table.rows[1:]:
                row.height = Cm(1.15)
                for cell in row.cells:
                    for run in cell.paragraphs[0].runs:
                        _set_font(run, size=15, bold=True)
    document.save(output_path)


def _render_broadcast_script(
    materials: list[FinalMaterial], participants: dict[str, Participant], output_path: Path
) -> None:
    blocks: list[str] = []
    for material in materials:
        participant_count = sum(len(heat) for heat in material.heats)
        lines = [
            f"下面进行{material.event_title}决赛，共{participant_count}名运动员。",
            f"比赛时间：{material.scheduled_time}。",
        ]
        position_word = "道" if material.position_header == "道次" else "位"
        for heat_index, assignments in enumerate(material.heats, start=1):
            if len(material.heats) > 1:
                lines.append(f"第{heat_index}组：")
            for participant_id, position in sorted(assignments, key=lambda item: item[1]):
                participant = participants[participant_id]
                lines.append(
                    f"第{position}{position_word}，{participant.bib}号，"
                    f"{participant.name}，{participant.unit}班。"
                )
        lines.append("请以上运动员提前到检录处检录，做好比赛准备。")
        blocks.append("\n".join(lines))
    output_path.write_text("\n\n".join(blocks) + "\n", encoding="utf-8-sig")


def _append_final_material(
    document: Document,
    material: FinalMaterial,
    participants: dict[str, Participant],
) -> None:
    _heading(document, f"{material.event_title} 决赛分组表")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(subtitle.add_run(f"比赛时间：{material.scheduled_time}"), size=11, bold=True)

    for heat_index, assignments in enumerate(material.heats, start=1):
        if heat_index > 1:
            document.add_paragraph()
        _heading(document, f"第{heat_index}组", size=12)
        rows = []
        for participant_id, lane in sorted(assignments, key=lambda item: item[1]):
            participant = participants[participant_id]
            rows.append([str(lane), participant.bib, participant.name, participant.unit, ""])
        _add_grid_table(
            document,
            [material.position_header, "号码", "姓名", "单位", "成绩"],
            rows,
        )


def render_result_ranking(
    event_title: str,
    ranked_results: list[tuple[int | None, Participant, Result]],
    output_path: str | Path,
) -> Path:
    document = Document()
    _configure_a4(document)
    _append_result_ranking(document, FinalRankingMaterial(event_title, ranked_results))
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


def render_result_rankings_batch(
    materials: list[FinalRankingMaterial], output_path: str | Path
) -> Path:
    if not materials:
        raise ValueError("没有可生成的决赛成绩排名")
    document = Document()
    _configure_a4(document)
    for index, material in enumerate(materials):
        if index:
            document.add_page_break()
        _append_result_ranking(document, material)
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


def _append_result_ranking(document: Document, material: FinalRankingMaterial) -> None:
    _heading(document, f"{material.event_title} 成绩排名表")
    status_names = {
        "VALID": "有效",
        "DNS": "未起跑",
        "DNF": "未完成",
        "DQ": "犯规",
        "NM": "无成绩",
    }
    rows = [
        [
            "" if rank is None else str(rank),
            participant.bib,
            participant.name,
            participant.unit,
            result.standard_display,
            status_names[result.status.value],
        ]
        for rank, participant, result in material.ranked_results
    ]
    _add_grid_table(document, ["名次", "号码", "姓名", "单位", "成绩", "状态"], rows)
    document.add_paragraph()
    signature = document.add_paragraph()
    signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(signature.add_run("裁判长签字：________________    日期：________________"), size=11)


def render_validation_report(report: ValidationReport, output_path: str | Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "核对报告"
    headers = ["规则代码", "级别", "具体问题", "对象类型", "对象ID", "来源文件", "工作表", "Excel行号", "原始值", "修正值"]
    sheet.append(headers)
    severity_names = {
        Severity.BLOCKING: "阻断",
        Severity.WARNING: "警告",
        Severity.AUTO_FIXED: "自动修正",
    }
    fills = {
        Severity.BLOCKING: PatternFill("solid", fgColor="FFC7CE"),
        Severity.WARNING: PatternFill("solid", fgColor="FFEB9C"),
        Severity.AUTO_FIXED: PatternFill("solid", fgColor="C6EFCE"),
    }
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for issue in report.issues:
        sheet.append(
            [
                issue.code, severity_names[issue.severity], issue.message, issue.entity_type,
                issue.entity_id,
                report.source_files.get(issue.source_file_id, issue.source_file_id[:8]),
                issue.source_sheet,
                issue.source_row,
                str(issue.original_value) if issue.original_value is not None else "",
                str(issue.normalized_value) if issue.normalized_value is not None else "",
            ]
        )
        for cell in sheet[sheet.max_row]:
            cell.fill = fills[issue.severity]
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [14, 10, 42, 16, 18, 12, 16, 10, 30, 30]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    workbook.close()
    return target
