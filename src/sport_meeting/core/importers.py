from __future__ import annotations

import datetime as dt
import hashlib
import math
import re
import uuid
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .models import (
    Athlete,
    Entry,
    EventRound,
    EventType,
    PerformanceKind,
    ProjectConfig,
    RoundType,
    Severity,
    SourceFileSnapshot,
    ValidationIssue,
    ValidationReport,
)


_ID_NAMESPACE = uuid.UUID("c96f29a1-5f42-4f06-b5b2-3d25dfa42c91")
_EVENT_ALIASES = {
    "一分钟跳绳": "1分钟跳绳",
    "1 分钟跳绳": "1分钟跳绳",
    "一分钟仰卧起坐": "1分钟仰卧起坐",
    "1 分钟仰卧起坐": "1分钟仰卧起坐",
    "4*100": "4×100米",
    "4x100": "4×100米",
    "4X100": "4×100米",
    "4 X 100米": "4×100米",
    "4×100": "4×100米",
    "4X100米接力": "4×100米接力",
    "4x100米接力": "4×100米接力",
    "4*100米接力": "4×100米接力",
}
_GRADE_NUMBER_CODES = {
    "初一": 1, "初二": 2, "初三": 3,
    "高一": 4, "高二": 5, "高三": 6,
}
_GRADE_ALIASES = {
    "初一": "初一", "初1": "初一", "初中一": "初一", "初中1": "初一",
    "七": "初一", "7": "初一", "初中七": "初一", "初中7": "初一",
    "初二": "初二", "初2": "初二", "初中二": "初二", "初中2": "初二",
    "八": "初二", "8": "初二", "初中八": "初二", "初中8": "初二",
    "初三": "初三", "初3": "初三", "初中三": "初三", "初中3": "初三",
    "九": "初三", "9": "初三", "初中九": "初三", "初中9": "初三",
    "高一": "高一", "高1": "高一", "高中一": "高一", "高中1": "高一",
    "十": "高一", "10": "高一", "高中十": "高一", "高中10": "高一",
    "高二": "高二", "高2": "高二", "高中二": "高二", "高中2": "高二",
    "十一": "高二", "11": "高二", "高中十一": "高二", "高中11": "高二",
    "高三": "高三", "高3": "高三", "高中三": "高三", "高中3": "高三",
    "十二": "高三", "12": "高三", "高中十二": "高三", "高中12": "高三",
}


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().replace("　", " ")


def normalize_event_name(value: object) -> tuple[str, bool]:
    original = _text(value)
    compact = re.sub(r"\s+", "", original).replace("X", "×").replace("x", "×").replace("*", "×")
    normalized = _EVENT_ALIASES.get(original, _EVENT_ALIASES.get(compact, compact))
    return normalized, normalized != original


def _allowed_registration_events(config: ProjectConfig, grade: str, sex: str) -> tuple[str, ...]:
    if grade.startswith("初"):
        return config.junior_boys_events if sex == "男" else config.junior_girls_events
    return config.senior_boys_events if sex == "男" else config.senior_girls_events


def _normalize_registration_event(
    value: object,
    grade: str,
    sex: str,
    config: ProjectConfig,
) -> tuple[str, bool, bool]:
    """Return normalized name, whether it changed, and whether it is configured.

    Registration cells are normalized with grade/sex context so an unambiguous
    shorthand can be corrected without turning a genuinely unknown event into a
    different competition silently.
    """
    original = _text(value)
    normalized, _changed = normalize_event_name(original)
    allowed = tuple(
        normalize_event_name(event)[0]
        for event in _allowed_registration_events(config, grade, sex)
        if _text(event)
    )

    candidate = normalized
    if re.fullmatch(r"\d{2,4}", candidate):
        metric_candidate = f"{candidate}米"
        if metric_candidate in allowed:
            candidate = metric_candidate
    elif candidate == "跳绳":
        matches = [event for event in allowed if "跳绳" in event and "集体" not in event]
        if len(matches) == 1:
            candidate = matches[0]
    elif candidate == "跳远":
        matches = [event for event in allowed if event in {"跳远", "立定跳远"}]
        if len(matches) == 1:
            candidate = matches[0]
    elif "仰卧起坐" in candidate:
        matches = [event for event in allowed if "仰卧起坐" in event]
        if len(matches) == 1:
            candidate = matches[0]

    return candidate, candidate != original, candidate in allowed


def _normalize_sex(value: object) -> str:
    text = _text(value)
    aliases = {"男生": "男", "男子": "男", "M": "男", "女生": "女", "女子": "女", "F": "女"}
    return aliases.get(text.upper(), aliases.get(text, text))


def _normalize_class(value: object) -> str:
    text = _text(value)
    match = re.search(r"\d+", text)
    return str(int(match.group())) if match else text.replace("班", "").strip("()（） ")


def _normalize_grade(value: object) -> str:
    text = _text(value).replace("年級", "年级").replace("級", "级")
    compact = re.sub(r"[\s()（）_\-—]+", "", text)
    compact = re.sub(r"(?:上|下)?学期$", "", compact)
    key = compact.removesuffix("年级")
    return _GRADE_ALIASES.get(key, text)


def _name_comparison_key(value: object) -> str:
    return re.sub(r"\s+", "", _text(value)).casefold()


def _chinese_integer(value: str) -> int | None:
    digits = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
    if value in digits:
        return digits[value]
    if value == "十":
        return 10
    if "十" not in value:
        return None
    tens_text, ones_text = value.split("十", 1)
    tens = digits.get(tens_text, 1) if tens_text else 1
    ones = digits.get(ones_text, 0) if ones_text else 0
    number = tens * 10 + ones
    return number if 1 <= number <= 99 else None


def _registration_scope_from_filename(path: Path) -> tuple[str, str]:
    compact = re.sub(r"\s+", "", path.stem).replace("年級", "年级").replace("級", "级")
    grade = ""
    grade_match = re.search(
        r"(初中?[一二三123]|初中?[七八九789](?:年级)?|"
        r"高中?[一二三123]|高中?(?:十二|十一|十|12|11|10)(?:年级)?|"
        r"[七八九789](?:年级)|(?:十二|十一|十|12|11|10)(?:年级))",
        compact,
    )
    if grade_match:
        grade = _normalize_grade(grade_match.group(1))
        if grade not in _GRADE_NUMBER_CODES:
            grade = ""

    class_name = ""
    class_match = re.search(r"[（(]?(\d{1,2})[)）]?班", compact)
    if class_match:
        number = int(class_match.group(1))
        if 1 <= number <= 99:
            class_name = str(number)
    else:
        chinese_class_match = re.search(r"([一二三四五六七八九十]{1,3})班", compact)
        if chinese_class_match:
            number = _chinese_integer(chinese_class_match.group(1))
            if number is not None:
                class_name = str(number)
    return grade, class_name


def _group_name(grade: str, sex: str) -> str:
    return f"{grade}{'男子' if sex == '男' else '女子'}组"


@dataclass(slots=True)
class RegistrationImport:
    athletes: list[Athlete] = field(default_factory=list)
    entries: list[Entry] = field(default_factory=list)
    report: ValidationReport = field(default_factory=ValidationReport)
    source_snapshots: list[SourceFileSnapshot] = field(default_factory=list)


@dataclass(slots=True)
class ScheduleImport:
    rounds: list[EventRound] = field(default_factory=list)
    report: ValidationReport = field(default_factory=ValidationReport)


def import_registration_workbooks(
    paths: Iterable[str | Path],
    config: ProjectConfig | None = None,
) -> RegistrationImport:
    result = RegistrationImport()
    effective_config = config or ProjectConfig()
    enforce_event_config = config is not None
    athletes_by_bib: dict[str, Athlete] = {}
    entry_keys: set[tuple[str, str, str]] = set()
    next_auto_sequence: dict[tuple[str, str, str], int] = {}
    first_name_occurrence: dict[tuple[str, str, str], tuple[str, str, int]] = {}
    auto_number_count = 0
    imported_at = dt.datetime.now().isoformat(timespec="seconds")

    for source_index, raw_path in enumerate(paths, start=1):
        path = Path(raw_path)
        source_id = uuid.uuid5(_ID_NAMESPACE, str(path.resolve()).lower()).hex
        result.report.source_files[source_id] = path.name
        try:
            content_hash = hashlib.sha256(path.read_bytes()).hexdigest()
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            result.report.add(
                ValidationIssue(
                    "IMPORT_001", Severity.BLOCKING, f"无法读取工作簿：{exc}",
                    source_file_id=source_id,
                )
            )
            continue

        filename_grade, filename_class = _registration_scope_from_filename(path)

        workbook_grades: set[str] = set()
        candidate_records: list[tuple[str, int, str, str]] = []
        for candidate_sheet in workbook.worksheets:
            if (
                candidate_sheet.sheet_state != "visible"
                or "不要导入" in candidate_sheet.title
                or candidate_sheet.title == "填写说明"
            ):
                continue
            candidate_header_row = None
            candidate_headers: list[str] = []
            for candidate_row_number, candidate_row in enumerate(
                candidate_sheet.iter_rows(
                    min_row=1, max_row=min(20, candidate_sheet.max_row), values_only=True
                ),
                start=1,
            ):
                header_values = [_text(value) for value in candidate_row]
                if "号码" in header_values and "姓名" in header_values:
                    candidate_header_row = candidate_row_number
                    candidate_headers = header_values
                    break
            if candidate_header_row is None:
                continue
            candidate_index = {
                name: position for position, name in enumerate(candidate_headers) if name
            }
            if not all(name in candidate_index for name in ("号码", "姓名", "年级", "班级")):
                continue
            for candidate_row_number, candidate_row in enumerate(
                candidate_sheet.iter_rows(
                    min_row=candidate_header_row + 1, values_only=True
                ),
                start=candidate_header_row + 1,
            ):
                candidate_values = list(candidate_row)
                candidate_name = _text(
                    candidate_values[candidate_index["姓名"]]
                    if candidate_index["姓名"] < len(candidate_values) else None
                )
                if not candidate_name:
                    continue
                candidate_grade = _normalize_grade(
                    candidate_values[candidate_index["年级"]]
                    if candidate_index["年级"] < len(candidate_values) else None
                )
                candidate_class = _normalize_class(
                    candidate_values[candidate_index["班级"]]
                    if candidate_index["班级"] < len(candidate_values) else None
                )
                if candidate_grade:
                    workbook_grades.add(candidate_grade)
                candidate_records.append(
                    (candidate_sheet.title, candidate_row_number, candidate_grade, candidate_class)
                )
        result.source_snapshots.append(
            SourceFileSnapshot(
                source_id,
                str(path.resolve()),
                content_hash,
                len(workbook.worksheets),
                len(candidate_records),
                imported_at,
            )
        )
        workbook_default_grade = filename_grade or (
            next(iter(workbook_grades)) if len(workbook_grades) == 1 else ""
        )
        class_corrections: dict[tuple[str, int], str] = {}
        classes_by_grade: dict[str, list[tuple[str, int, str]]] = {}
        for sheet_title, row_number, candidate_grade, candidate_class in candidate_records:
            effective_grade = candidate_grade or workbook_default_grade
            if effective_grade and candidate_class:
                classes_by_grade.setdefault(effective_grade, []).append(
                    (sheet_title, row_number, candidate_class)
                )
        for grade_records in classes_by_grade.values():
            if filename_grade and not filename_class:
                continue
            if filename_class:
                continue
            counts = Counter(class_name for _sheet, _row, class_name in grade_records)
            if len(counts) != 2:
                continue
            (dominant_class, dominant_count), (minority_class, minority_count) = counts.most_common()
            total_count = dominant_count + minority_count
            if dominant_count < 4 or minority_count != 1 or dominant_count / total_count < 0.8:
                continue
            for sheet_title, row_number, class_name in grade_records:
                if class_name == minority_class:
                    class_corrections[(sheet_title, row_number)] = dominant_class

        for grade, grade_records in classes_by_grade.items():
            if filename_grade or filename_class:
                continue
            effective_records = [
                (
                    sheet_title,
                    row_number,
                    class_corrections.get((sheet_title, row_number), class_name),
                )
                for sheet_title, row_number, class_name in grade_records
            ]
            effective_counts = Counter(class_name for _sheet, _row, class_name in effective_records)
            if len(effective_counts) <= 1:
                continue
            distribution = "、".join(
                f"{class_name}班 {count} 人"
                for class_name, count in sorted(
                    effective_counts.items(),
                    key=lambda item: (int(item[0]) if item[0].isdigit() else 999, item[0]),
                )
            )
            result.report.add(
                ValidationIssue(
                    "IMPORT_019", Severity.WARNING,
                    f"同一报名文件的{grade}出现多个班级值（{distribution}），"
                    "无法安全自动修正，请人工核对",
                    entity_type="CLASS", entity_id=grade, source_file_id=source_id,
                    source_sheet="跨工作表汇总",
                    original_value=[
                        {"工作表": sheet_title, "Excel行号": row_number, "班级": class_name}
                        for sheet_title, row_number, class_name in effective_records
                    ],
                    normalized_value={"班级分布": dict(effective_counts)},
                )
            )

        workbook_classes = (
            {filename_class}
            if filename_class
            else {
                class_corrections.get((sheet_title, row_number), candidate_class)
                for sheet_title, row_number, _grade, candidate_class in candidate_records
                if candidate_class
            }
        )
        workbook_default_class = next(iter(workbook_classes)) if len(workbook_classes) == 1 else ""

        for worksheet in workbook.worksheets:
            if (
                worksheet.sheet_state != "visible"
                or "不要导入" in worksheet.title
                or worksheet.title == "填写说明"
            ):
                continue
            header_row = None
            headers: list[str] = []
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=1, max_row=min(20, worksheet.max_row), values_only=True), start=1
            ):
                candidate = [_text(value) for value in row]
                if "号码" in candidate and "姓名" in candidate:
                    header_row, headers = row_number, candidate
                    break
            if header_row is None:
                result.report.add(
                    ValidationIssue(
                        "IMPORT_002", Severity.WARNING, "未找到包含“号码、姓名”的表头，已跳过 Sheet",
                        source_file_id=source_id, source_sheet=worksheet.title,
                    )
                )
                continue

            index = {name: position for position, name in enumerate(headers) if name}
            required = [name for name in ("号码", "姓名", "性别", "年级", "班级") if name not in index]
            if required:
                result.report.add(
                    ValidationIssue(
                        "IMPORT_003", Severity.BLOCKING, f"报名表缺少字段：{'、'.join(required)}",
                        source_file_id=source_id, source_sheet=worksheet.title, source_row=header_row,
                    )
                )
                continue
            event_columns = [position for position, name in enumerate(headers) if name.startswith("项目")]

            data_rows = list(
                enumerate(
                    worksheet.iter_rows(min_row=header_row + 1, values_only=True),
                    start=header_row + 1,
                )
            )
            sheet_grades: set[str] = set()
            sheet_classes: set[str] = set()
            for _row_number, source_row in data_rows:
                source_values = list(source_row)
                source_name = _text(
                    source_values[index["姓名"]] if index["姓名"] < len(source_values) else None
                )
                if not source_name:
                    continue
                source_grade = _normalize_grade(
                    source_values[index["年级"]] if index["年级"] < len(source_values) else None
                )
                source_class = filename_class or class_corrections.get(
                    (worksheet.title, _row_number),
                    _normalize_class(
                        source_values[index["班级"]]
                        if index["班级"] < len(source_values) else None
                    ),
                )
                if source_grade:
                    sheet_grades.add(source_grade)
                if source_class:
                    sheet_classes.add(source_class)
            sheet_default_grade = next(iter(sheet_grades)) if len(sheet_grades) == 1 else ""
            sheet_default_class = next(iter(sheet_classes)) if len(sheet_classes) == 1 else ""
            last_grade = ""
            last_class = ""

            for row_number, row in data_rows:
                values = list(row)
                provided_bib = _text(values[index["号码"]] if index["号码"] < len(values) else None)
                name = _text(values[index["姓名"]] if index["姓名"] < len(values) else None)
                if not provided_bib and not name:
                    continue
                if provided_bib == "号码" or name == "姓名":
                    continue
                sex = _normalize_sex(values[index["性别"]] if index["性别"] < len(values) else None)
                raw_grade = _text(values[index["年级"]] if index["年级"] < len(values) else None)
                normalized_raw_grade = _normalize_grade(raw_grade)
                raw_class_text = _text(values[index["班级"]] if index["班级"] < len(values) else None)
                raw_class = _normalize_class(raw_class_text)
                corrected_raw_class = filename_class or class_corrections.get(
                    (worksheet.title, row_number), raw_class
                )
                grade = filename_grade or normalized_raw_grade or last_grade or sheet_default_grade or workbook_default_grade
                class_name = corrected_raw_class or last_class or sheet_default_class or workbook_default_class
                filled_fields: list[str] = []
                if not raw_grade and grade:
                    filled_fields.append("年级")
                if not raw_class and class_name:
                    filled_fields.append("班级")
                if grade:
                    last_grade = grade
                if class_name:
                    last_class = class_name
                if filled_fields:
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_015", Severity.AUTO_FIXED,
                            f"{'、'.join(filled_fields)}为空，已按同一报名表内容自动补全",
                            entity_type="ATHLETE", entity_id=name, source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_number,
                            original_value={"年级": raw_grade or None, "班级": raw_class_text or None},
                            normalized_value={"年级": grade, "班级": class_name},
                        )
                    )
                filename_overrides: dict[str, dict[str, str | None]] = {}
                if filename_grade and raw_grade and normalized_raw_grade != filename_grade:
                    filename_overrides["年级"] = {"原值": raw_grade, "文件名识别值": filename_grade}
                if filename_class and raw_class and raw_class != filename_class:
                    filename_overrides["班级"] = {"原值": raw_class_text, "文件名识别值": filename_class}
                if filename_overrides:
                    scope_text = filename_grade
                    if filename_class:
                        scope_text += f"({filename_class})班"
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_020", Severity.AUTO_FIXED,
                            f"报名文件名已明确识别为{scope_text}，已以文件名为准修正",
                            entity_type="ATHLETE", entity_id=name, source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_number,
                            original_value=filename_overrides,
                            normalized_value={"年级": grade, "班级": class_name},
                        )
                    )
                if not filename_class and raw_class and corrected_raw_class != raw_class:
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_018", Severity.AUTO_FIXED,
                            "同一报名表内该班级值与其余学生明显不一致，"
                            "已按同年级唯一主班级自动修正",
                            entity_type="ATHLETE", entity_id=name, source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_number,
                            original_value=raw_class_text, normalized_value=corrected_raw_class,
                        )
                    )
                if raw_grade and normalized_raw_grade and normalized_raw_grade != raw_grade:
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_016", Severity.AUTO_FIXED,
                            "年级表述已统一为系统标准写法",
                            entity_type="ATHLETE", entity_id=name, source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_number,
                            original_value=raw_grade, normalized_value=normalized_raw_grade,
                        )
                    )
                if not all((name, sex, grade, class_name)):
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_004", Severity.BLOCKING, "运动员必要字段不完整",
                            entity_type="ATHLETE", entity_id=provided_bib, source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_number,
                        )
                    )
                    continue
                grade_code = _GRADE_NUMBER_CODES.get(grade)
                if grade_code is None or not class_name.isdigit() or sex not in {"男", "女"}:
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_011", Severity.BLOCKING,
                            "自动编号要求年级为初一至高三、班级为数字、性别为男或女",
                            entity_type="ATHLETE", source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_number,
                        )
                    )
                    continue
                class_number = int(class_name)
                if not 1 <= class_number <= 99:
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_011", Severity.BLOCKING, "自动编号要求班级号为 1 至 99",
                            entity_type="ATHLETE", source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_number,
                        )
                    )
                    continue
                normalized_events: list[str] = []
                for column in event_columns:
                    if column >= len(values) or values[column] in (None, ""):
                        continue
                    event_name, changed, configured = _normalize_registration_event(
                        values[column], grade, sex, effective_config
                    )
                    if not event_name or event_name in {"项目1", "项目2", "项目3"}:
                        continue
                    if changed:
                        result.report.add(
                            ValidationIssue(
                                "IMPORT_007", Severity.AUTO_FIXED, "项目名称已按年级、性别和项目设置归一化",
                                entity_type="ATHLETE", entity_id=name, source_file_id=source_id,
                                source_sheet=worksheet.title, source_row=row_number,
                                original_value=_text(values[column]), normalized_value=event_name,
                            )
                        )
                    if not configured and enforce_event_config:
                        result.report.add(
                            ValidationIssue(
                                "IMPORT_021", Severity.BLOCKING,
                                f"项目“{event_name}”不在{grade}{'男生' if sex == '男' else '女生'}项目设置中",
                                entity_type="ATHLETE", entity_id=name, source_file_id=source_id,
                                source_sheet=worksheet.title, source_row=row_number,
                                original_value=_text(values[column]), normalized_value=event_name,
                            )
                        )
                        continue
                    if event_name not in normalized_events:
                        normalized_events.append(event_name)
                if not normalized_events:
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_022", Severity.WARNING,
                            "填写了姓名但没有有效报名项目，已从有效运动员名单中排除",
                            entity_type="ATHLETE", entity_id=name, source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_number,
                            original_value=[
                                _text(values[column])
                                for column in event_columns
                                if column < len(values) and values[column] not in (None, "")
                            ],
                        )
                    )
                    continue
                duplicate_key = (grade, class_name, _name_comparison_key(name))
                first_occurrence = first_name_occurrence.get(duplicate_key)
                if first_occurrence is None:
                    first_name_occurrence[duplicate_key] = (path.name, worksheet.title, row_number)
                else:
                    first_file, first_sheet, first_row = first_occurrence
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_017", Severity.WARNING,
                            f"同一班级内姓名“{name}”重复，可能是重复登记；"
                            f"首次出现于 {first_file} / {first_sheet} / 第 {first_row} 行",
                            entity_type="ATHLETE", entity_id=f"{grade}({class_name})班/{name}",
                            source_file_id=source_id, source_sheet=worksheet.title,
                            source_row=row_number,
                            original_value={"年级": grade, "班级": class_name, "姓名": name},
                            normalized_value={
                                "首次出现文件": first_file,
                                "首次出现工作表": first_sheet,
                                "首次出现行": first_row,
                            },
                        )
                    )
                counter_key = (grade, class_name, sex)
                sequence = next_auto_sequence.get(counter_key, 1 if sex == "男" else 2)
                while f"{grade_code}{class_number:02d}{sequence:02d}" in athletes_by_bib:
                    sequence += 2
                if sequence > 99:
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_011", Severity.BLOCKING, "同一班级同性别运动员过多，无法生成两位班内序号",
                            entity_type="ATHLETE", source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_number,
                        )
                    )
                    continue
                bib = f"{grade_code}{class_number:02d}{sequence:02d}"
                next_auto_sequence[counter_key] = sequence + 2
                auto_number_count += 1
                if provided_bib:
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_014", Severity.AUTO_FIXED,
                            "报名表中的自编号码已忽略，已按系统规则重新编号",
                            entity_type="ATHLETE", entity_id=bib, source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_number,
                            original_value=provided_bib, normalized_value=bib,
                        )
                    )
                athlete = Athlete(uuid.uuid5(_ID_NAMESPACE, f"athlete:{bib}").hex, bib, name, sex, grade, class_name)
                existing = athletes_by_bib.get(bib)
                if existing and existing != athlete:
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_005", Severity.BLOCKING, "全校运动员号码重复且资料不一致",
                            entity_type="ATHLETE", entity_id=bib, source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_number,
                            original_value={"姓名": existing.name, "性别": existing.sex, "年级": existing.grade, "班级": existing.class_name},
                            normalized_value={"姓名": name, "性别": sex, "年级": grade, "班级": class_name},
                        )
                    )
                    continue
                if not existing:
                    athletes_by_bib[bib] = athlete
                else:
                    result.report.add(
                        ValidationIssue(
                            "IMPORT_006", Severity.AUTO_FIXED, "重复的相同运动员资料已合并",
                            entity_type="ATHLETE", entity_id=bib, source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_number,
                        )
                    )

                group = _group_name(grade, sex)
                for event_name in normalized_events:
                    key = (athlete.id, group, event_name)
                    if key in entry_keys:
                        continue
                    entry_keys.add(key)
                    result.entries.append(
                        Entry(
                            uuid.uuid5(_ID_NAMESPACE, f"entry:{athlete.id}:{group}:{event_name}").hex,
                            athlete.id, event_name, group, source_id, worksheet.title, row_number,
                        )
                    )
        workbook.close()
    if auto_number_count:
        result.report.add(
            ValidationIssue(
                "IMPORT_012", Severity.AUTO_FIXED,
                f"已按系统规则为 {auto_number_count} 名运动员统一编号",
                entity_type="ATHLETE", source_file_id="__aggregate__",
                source_sheet="跨文件汇总",
                normalized_value=auto_number_count,
            )
        )
    entries_by_athlete: dict[str, list[Entry]] = {}
    for entry in result.entries:
        entries_by_athlete.setdefault(entry.athlete_id, []).append(entry)
    for athlete_id, athlete_entries in entries_by_athlete.items():
        if len(athlete_entries) <= 2:
            continue
        athlete = next(item for item in athletes_by_bib.values() if item.id == athlete_id)
        result.report.add(
            ValidationIssue(
                "IMPORT_010",
                Severity.BLOCKING,
                f"每名运动员最多报名两个项目，当前报名 {len(athlete_entries)} 项",
                entity_type="ATHLETE",
                entity_id=athlete.bib,
                source_file_id=athlete_entries[0].source_file_id,
                source_sheet=athlete_entries[0].source_sheet,
                source_row=athlete_entries[0].source_row,
                original_value=[item.event_name for item in athlete_entries],
            )
        )
    classes_by_grade: dict[str, set[int]] = {}
    for athlete in athletes_by_bib.values():
        if athlete.class_name.isdigit():
            classes_by_grade.setdefault(athlete.grade, set()).add(int(athlete.class_name))
    result.report.source_files["__aggregate__"] = "批量汇总检查"
    for grade, class_numbers in sorted(
        classes_by_grade.items(), key=lambda item: _GRADE_NUMBER_CODES.get(item[0], 99)
    ):
        if len(class_numbers) < 2:
            continue
        first_class, last_class = min(class_numbers), max(class_numbers)
        missing = sorted(set(range(first_class, last_class + 1)) - class_numbers)
        if missing:
            missing_text = "、".join(f"{grade}({number})班" for number in missing)
            result.report.add(
                ValidationIssue(
                    "IMPORT_013", Severity.WARNING,
                    f"可能漏交班级报名表：{missing_text}；已导入班号范围为 {first_class} 至 {last_class}",
                    entity_type="CLASS", entity_id=grade, source_file_id="__aggregate__",
                    source_sheet="跨文件汇总", original_value=sorted(class_numbers),
                    normalized_value=missing,
                )
            )
    result.athletes = sorted(athletes_by_bib.values(), key=lambda item: item.bib)
    return result


def _round_type(name: str) -> RoundType:
    if "预决赛" in name:
        return RoundType.TIMED_FINAL
    if "决赛" in name:
        return RoundType.FINAL
    if "预赛" in name:
        return RoundType.PRELIMINARY
    return RoundType.TIMED_FINAL


def _event_type(value: str, event_name: str) -> EventType:
    if any(word in event_name for word in ("接力", "集体", "拔河", "袋鼠跳")):
        return EventType.TEAM
    if any(word in event_name for word in ("跳远", "跳高", "跳绳", "仰卧起坐", "引体向上", "铅球", "实心球")):
        return EventType.FIELD
    if "田" in value:
        return EventType.FIELD
    if "集体" in value:
        return EventType.TEAM
    return EventType.TRACK


def _format_time(value: object) -> str:
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, dt.time):
        return value.strftime("%H:%M")
    return _text(value)


def parse_available_lanes(value: object) -> tuple[int, ...]:
    text = _text(value)
    if not text:
        return ()
    text = text.replace("，", ",").replace("—", "-").replace("~", "-")
    lanes: list[int] = []
    for part in text.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start, end = (int(item.strip()) for item in part.split("-", 1))
            lanes.extend(range(start, end + 1))
        else:
            lanes.append(int(part))
    return tuple(lanes)


def _integer(value: object) -> int | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    match = re.search(r"\d+", _text(value))
    return int(match.group()) if match else None


def _performance_for(event_name: str, event_type: EventType) -> tuple[PerformanceKind, str, str]:
    if event_type is EventType.TRACK or "接力" in event_name:
        distance_match = re.search(r"(\d+)\s*米", event_name)
        default_unit = "MINUTE" if distance_match and int(distance_match.group(1)) >= 800 else "SECOND"
        return PerformanceKind.TIME, "MILLISECOND", default_unit
    if any(word in event_name for word in ("跳绳", "仰卧起坐", "引体向上")):
        return PerformanceKind.COUNT, "COUNT", "COUNT"
    if any(word in event_name for word in ("跳远", "跳高", "铅球", "实心球")):
        return PerformanceKind.DISTANCE, "MILLIMETER", "METER"
    return PerformanceKind.MANUAL_RANK, "", ""


def _strip_round_suffix(value: str) -> str:
    return re.sub(r"(?:预决赛|预赛|决赛)$", "", value).strip()


def _infer_group_and_event(full_name: str, group_name: str, raw_event: str) -> tuple[str, str]:
    group = group_name.replace("组组", "组").strip()
    if raw_event:
        event, _ = normalize_event_name(_strip_round_suffix(raw_event))
        return group, event

    base = _strip_round_suffix(full_name)
    school_group = re.match(r"^(初一|初二|初三|高一|高二)(?:年级)?(?:(男子|女子))?组?", base)
    if school_group:
        grade, sex = school_group.groups()
        if sex:
            group = f"{grade}{sex}组"
        elif not group:
            group = f"{grade}年级组"
        base = base[school_group.end():]
    elif base.startswith("教工组"):
        group = "教工组"
        base = base[len("教工组"):]
    elif group:
        prefixes = {group, group.removesuffix("组"), group.replace("年级组", "年级")}
        for prefix in sorted(prefixes, key=len, reverse=True):
            if prefix and base.startswith(prefix):
                base = base[len(prefix):]
                break

    event = re.sub(r"(?:比赛|赛)$", "", base).strip()
    event, _ = normalize_event_name(event)
    return group, event


def _heat_count(
    explicit: int | None,
    round_type: RoundType,
    event_type: EventType,
    event_name: str,
    declared: int | None,
    available_lanes: tuple[int, ...],
) -> tuple[int, bool]:
    if explicit:
        return explicit, False
    if round_type is RoundType.FINAL or event_type is EventType.FIELD:
        return 1, True
    if event_type is EventType.TEAM:
        return 1, True
    if declared:
        distance_match = re.search(r"(?<!×)(\d+)\s*米", event_name)
        distance = int(distance_match.group(1)) if distance_match else None
        capacity = len(available_lanes) or 8
        if distance in {50, 100, 200}:
            return max(1, math.ceil(declared / capacity)), True
        if distance == 400:
            return max(1, math.ceil(declared / 16)), True
        if distance and distance >= 800:
            return 1, True
    return 0, False


def import_schedule_workbook(path: str | Path, config=None) -> ScheduleImport:
    """导入结构化日程；兼容老系统的“竞赛项目/人数/类型/组别/项目”字段。"""
    result = ScheduleImport()
    source = Path(path)
    source_id = uuid.uuid5(_ID_NAMESPACE, str(source.resolve()).lower()).hex
    result.report.source_files[source_id] = source.name
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        result.report.add(ValidationIssue("SCHEDULE_003", Severity.BLOCKING, f"无法读取日程：{exc}", source_file_id=source_id))
        return result
    try:
        for worksheet in workbook.worksheets:
            if worksheet.title == "填写说明" or "不要导入" in worksheet.title:
                continue
            header_row = None
            headers: list[str] = []
            for row_no, row in enumerate(worksheet.iter_rows(min_row=1, max_row=min(20, worksheet.max_row), values_only=True), 1):
                candidate = [_text(value) for value in row]
                if any(name in candidate for name in ("竞赛项目", "项目名称", "项目")) and any(
                    name in candidate for name in ("人数", "总人数/队数", "比赛时间", "时间")
                ):
                    header_row, headers = row_no, candidate
                    break
            if header_row is None:
                result.report.add(
                    ValidationIssue("SCHEDULE_004", Severity.WARNING, "未找到结构化日程表头，已跳过 Sheet", source_file_id=source_id, source_sheet=worksheet.title)
                )
                continue
            indexes = {name: i for i, name in enumerate(headers) if name}

            def cell(values, *names):
                for name in names:
                    position = indexes.get(name)
                    if position is not None and position < len(values):
                        return values[position]
                return None

            schedule_rows: list[tuple[str, str, RoundType, int]] = []
            for row_no, row in enumerate(worksheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                values = list(row)
                full_name = _text(cell(values, "竞赛项目", "项目名称"))
                raw_event = _text(cell(values, "项目"))
                group_name = _text(cell(values, "组别"))
                if not full_name and not raw_event:
                    continue
                row_source = dict(source_file_id=source_id, source_sheet=worksheet.title, source_row=row_no)
                required_cells = (
                    (("日期",), "日期"),
                    (("时间段", "上下午"), "时间段"),
                    (("序号",), "序号"),
                )
                for aliases, label in required_cells:
                    if any(alias in indexes for alias in aliases) and not _text(cell(values, *aliases)):
                        result.report.add(ValidationIssue(
                            "SCHEDULE_009", Severity.BLOCKING, f"{label}不能为空", **row_source,
                        ))
                round_type = _round_type(full_name or raw_event)
                title_group, _ = _infer_group_and_event(full_name, "", "") if full_name else ("", "")
                structured_group = group_name
                group_name, event_name = _infer_group_and_event(full_name, group_name, raw_event)
                title_group_specific = any(sex in title_group for sex in ("男子", "女子"))
                structured_group_specific = any(sex in structured_group for sex in ("男子", "女子"))
                if (
                    title_group and structured_group and title_group != structured_group
                    and title_group_specific and structured_group_specific
                ):
                    result.report.add(ValidationIssue(
                        "SCHEDULE_010", Severity.BLOCKING,
                        f"项目名称中的组别“{title_group}”与结构化组别“{structured_group}”不一致",
                        original_value=title_group, normalized_value=structured_group, **row_source,
                    ))
                if not group_name or not event_name:
                    result.report.add(
                        ValidationIssue(
                            "SCHEDULE_005", Severity.BLOCKING, "无法从日程行确定组别或项目，请补充结构化字段",
                            source_file_id=source_id, source_sheet=worksheet.title, source_row=row_no,
                            original_value=full_name or raw_event,
                        )
                    )
                    continue
                type_text = _text(cell(values, "比赛类型", "类型"))
                event_type = _event_type(type_text, event_name)
                declared = _integer(cell(values, "人数", "总人数/队数"))
                try:
                    available_lanes = parse_available_lanes(cell(values, "可用道次"))
                except Exception:
                    result.report.add(
                        ValidationIssue(
                            "SCHEDULE_007", Severity.BLOCKING, "可用道次格式错误",
                            source_file_id=source_id, source_sheet=worksheet.title, source_row=row_no,
                        )
                    )
                    continue
                heat_count, inferred_heat_count = _heat_count(
                    _integer(cell(values, "组数", "总组数")), round_type, event_type,
                    event_name, declared, available_lanes
                )
                if inferred_heat_count:
                    result.report.add(
                        ValidationIssue(
                            "SCHEDULE_008", Severity.AUTO_FIXED, f"日程未提供组数，已建议为 {heat_count} 组",
                            entity_type="EVENT_ROUND", source_file_id=source_id,
                            source_sheet=worksheet.title, source_row=row_no,
                            original_value=None, normalized_value=heat_count,
                        )
                    )
                elif heat_count == 0:
                    result.report.add(
                        ValidationIssue(
                            "SCHEDULE_006", Severity.WARNING, "日程未提供组数，需在编排前确认",
                            source_file_id=source_id, source_sheet=worksheet.title, source_row=row_no,
                        )
                    )
                date_text = _format_time(cell(values, "日期"))
                period = _text(cell(values, "时间段", "上下午"))
                time_text = _format_time(cell(values, "比赛时间", "时间"))
                if config and config.meeting_start_date and config.meeting_end_date and date_text:
                    try:
                        date_value = dt.date.fromisoformat(date_text[:10]) if re.match(r"^\d{4}-\d{2}-\d{2}", date_text) else None
                    except ValueError:
                        date_value = None
                    if date_value and not (dt.date.fromisoformat(config.meeting_start_date) <= date_value <= dt.date.fromisoformat(config.meeting_end_date)):
                        result.report.add(ValidationIssue(
                            "SCHEDULE_011", Severity.BLOCKING,
                            f"比赛日期 {date_value.isoformat()} 超出项目配置的运动会日期范围",
                            original_value=date_text, normalized_value=f"{config.meeting_start_date} 至 {config.meeting_end_date}", **row_source,
                        ))
                scheduled = " ".join(part for part in (date_text, period, time_text) if part)
                kind, canonical_unit, default_unit = _performance_for(event_name, event_type)
                event_id = uuid.uuid5(
                    _ID_NAMESPACE,
                    f"round:{group_name}:{event_name}:{round_type.value}:{scheduled}:{row_no}",
                ).hex
                result.rounds.append(
                    EventRound(
                        event_id, event_name, group_name, event_type, round_type, heat_count,
                        scheduled, declared, available_lanes, kind, canonical_unit, default_unit,
                    )
                )
                schedule_rows.append((group_name, event_name, round_type, row_no))
                distance_match = re.search(r"(?<!×)(\d+)\s*米", event_name)
                distance = int(distance_match.group(1)) if distance_match else None
                if event_type is EventType.TRACK and declared and distance in {50, 100, 200, 400}:
                    capacity = len(available_lanes) or (config.track_lanes if config else 8)
                    max_per_heat = 16 if distance == 400 else capacity
                    needed = math.ceil(declared / max_per_heat)
                    if heat_count < needed:
                        result.report.add(ValidationIssue(
                            "SCHEDULE_012", Severity.BLOCKING if distance != 400 else Severity.WARNING,
                            f"{event_name}共{declared}人、{heat_count}组，按每組最多{max_per_heat}人至少需要{needed}组",
                            **row_source,
                        ))
            counts: dict[tuple[str, str], list[tuple[RoundType, int]]] = {}
            for group, event, round_kind, source_row in schedule_rows:
                counts.setdefault((group, event), []).append((round_kind, source_row))
            if config:
                from .schedule_planner import guess_event_round_mode
                for (group, event), rounds_for_event in counts.items():
                    if group == "教工组":
                        continue
                    grade = re.match(r"^(初一|初二|初三|高一|高二|高三)", group)
                    sex = "boys" if "男子" in group else "girls" if "女子" in group else ""
                    if not grade or not sex:
                        continue
                    level = "junior" if grade.group(1).startswith("初") else "senior"
                    formats = getattr(config, f"{level}_{sex}_event_formats")
                    mode = formats.get(event, guess_event_round_mode(event))
                    if mode == "PRELIMINARY_FINAL":
                        for expected, label in ((RoundType.PRELIMINARY, "预赛"), (RoundType.FINAL, "决赛")):
                            matching = [(kind, source_row) for kind, source_row in rounds_for_event if kind is expected]
                            if len(matching) != 1:
                                source_row = (matching[1][1] if len(matching) > 1 else rounds_for_event[0][1])
                                result.report.add(ValidationIssue(
                                    "SCHEDULE_013", Severity.BLOCKING,
                                    f"配置为“预赛＋决赛”的项目必须恰有一场{label}，当前{len(matching)}场",
                                    source_file_id=source_id, source_sheet=worksheet.title, source_row=source_row,
                                ))
    finally:
        workbook.close()
    return result
