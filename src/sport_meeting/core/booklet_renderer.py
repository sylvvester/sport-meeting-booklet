from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field, replace
from pathlib import Path
from datetime import date
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from .models import Athlete, EventRound, EventType, HeatAssignment, ProjectConfig, RoundType
from .word_layout import fit_single_line_font, normalize_single_line_text, prevent_cell_text_wrap


@dataclass(frozen=True, slots=True)
class RenderStats:
    event_count: int
    table_count: int
    assignment_count: int
    font_adjustment_count: int = 0
    layout_warnings: tuple[str, ...] = ()


@dataclass(slots=True)
class _LayoutDiagnostics:
    font_adjustment_count: int = 0
    warnings: list[str] = field(default_factory=list)


def render_full_booklet(
    path: str | Path,
    rounds: list[EventRound],
    assignments: list[HeatAssignment],
    athletes: list[Athlete],
    title: str | None = None,
    config: ProjectConfig | None = None,
    include_cover: bool = True,
) -> RenderStats:
    """生成可连续打印的整册，动态写入日程、代表队名单和竞赛分组。

    include_cover=True（默认）时输出标题封面；封面本身不标页码，页码从目录页
    开始按 1 编号（通过 pgNumType/@w:start 重启编号）。
    """
    config = config or ProjectConfig()
    title = title or config.booklet_title
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _configure_document(document)
    diagnostics = _LayoutDiagnostics()

    if include_cover:
        cover = document.add_paragraph(style="Title")
        cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cover.paragraph_format.space_before = Pt(170)
        _set_run_font(cover.add_run(title), "微软雅黑", 28, bold=True)
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_before = Pt(36)
        _set_run_font(subtitle.add_run("竞赛秩序册"), "微软雅黑", 20, bold=True)
        # 封面单独成一节：本节首页页脚留空（不标页码），
        # 下一节从目录页把页码重新按 1 开始。
        _set_blank_run_font(document.sections[0].first_page_footer.paragraphs[0])
        document.add_section(WD_SECTION.NEW_PAGE)
        # add_section 会把这组设置重置，需在新节建立后重新指定。
        document.sections[0].different_first_page_header_footer = True
        _restart_page_numbering(document)

    _add_chapter_title(document, "目录", size=36, page_break=include_cover)
    contents = (
        "一、运动会开幕式议程", "二、运动会组委会名单", "三、运动会仲裁委员会及裁判员名单",
        "四、运动会竞赛规程", "五、运动会竞赛日程", "六、运动员代表队名单",
        "七、运动会竞赛分组表", "八、示例学校田径运动会校记录",
    )
    for item in contents:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(12)
        _set_run_font(paragraph.add_run(item), "微软雅黑", 15, bold=True)

    static_chapters = (
        ("一、开幕式议程", config.opening_ceremony_text, "请在软件中编辑本届运动会开幕式议程。"),
        ("二、组委会名单", config.organizing_committee_text, "请在软件中编辑本届运动会组委会名单。"),
        ("三、仲裁委员会及裁判员名单", config.officials_text, "请在软件中编辑仲裁委员会及裁判员名单。"),
        ("四、竞赛规程", config.competition_rules_text, "请在软件中编辑本届运动会竞赛规程。"),
    )
    for heading, content, placeholder in static_chapters:
        _add_page_break(document)
        _add_chapter_title(document, heading)
        _append_static_chapter(document, content, placeholder, heading)

    _add_page_break(document)
    _append_schedule(document, rounds, diagnostics)
    _add_page_break(document)
    _append_team_rosters(document, athletes, diagnostics)
    _add_page_break(document)
    stats = _append_competition_groups(document, rounds, assignments, athletes, diagnostics)
    _add_page_break(document)
    _add_chapter_title(document, "八、示例学校田径运动会校记录")
    _append_multiline_text(
        document,
        config.school_records_text,
        "请在软件中编辑最新校运动会纪录。",
        diagnostics=diagnostics,
        table_context="校纪录",
    )
    document.save(target)
    return replace(
        stats,
        font_adjustment_count=diagnostics.font_adjustment_count,
        layout_warnings=tuple(diagnostics.warnings),
    )


def render_competition_groups(
    path: str | Path,
    rounds: list[EventRound],
    assignments: list[HeatAssignment],
    athletes: list[Athlete],
) -> RenderStats:
    """生成与样例秩序册一致的五列表格竞赛分组文档。"""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _configure_document(document)
    diagnostics = _LayoutDiagnostics()
    stats = _append_competition_groups(document, rounds, assignments, athletes, diagnostics)
    document.save(target)
    return replace(
        stats,
        font_adjustment_count=diagnostics.font_adjustment_count,
        layout_warnings=tuple(diagnostics.warnings),
    )


def _session_sort_key(value: str) -> tuple[str, int]:
    """按“日期 → 上午/下午”排序，用于日程与分组表统一分块。"""
    day, session, _clock = _schedule_parts(value)
    return day, 0 if session == "上午" else 1


def _session_label(day: str, session: str) -> str:
    """把“2026-10-20 上午”统一渲染为“10月20日（星期二） 上午”。"""
    label = day
    try:
        parsed = date.fromisoformat(day)
        label = f"{parsed.month}月{parsed.day}日（星期{'一二三四五六日'[parsed.weekday()]}）"
    except ValueError:
        pass
    return f"{label} {session}".strip()


def _append_competition_groups(
    document: Document,
    rounds: list[EventRound],
    assignments: list[HeatAssignment],
    athletes: list[Athlete],
    diagnostics: _LayoutDiagnostics | None = None,
) -> RenderStats:
    _add_chapter_title(document, "七、竞赛分组表")

    athlete_map = {athlete.id: athlete for athlete in athletes}
    assignments_by_round: dict[str, list[HeatAssignment]] = defaultdict(list)
    for assignment in assignments:
        assignments_by_round[assignment.event_round_id].append(assignment)

    # 与第五章一致：先按“日期 → 上午/下午 → 径赛/田赛/集体项目”分块，
    # 各类型内部按比赛时间排列，并输出“径赛/田赛/集体项目”副标题。
    type_order = {EventType.TRACK: 0, EventType.FIELD: 1, EventType.TEAM: 2}
    type_names = {EventType.TRACK: "径赛", EventType.FIELD: "田赛", EventType.TEAM: "集体项目"}
    ordered_rounds = sorted(
        rounds,
        key=lambda item: (
            *_session_sort_key(item.scheduled_time),
            type_order[item.event_type],
            item.scheduled_time,
            item.group_name,
            item.event_name,
            item.round_type.value,
        ),
    )
    table_count = 0
    previous_session = None
    previous_type = None
    for sequence, round_ in enumerate(ordered_rounds, 1):
        event_assignments = assignments_by_round.get(round_.id, [])
        day, session, _clock = _schedule_parts(round_.scheduled_time)
        if (day, session) != previous_session:
            session_title = document.add_paragraph()
            session_title.paragraph_format.keep_with_next = True
            session_title.paragraph_format.space_before = Pt(10)
            session_title.paragraph_format.space_after = Pt(6)
            _set_run_font(session_title.add_run(_session_label(day, session)), "黑体", 13, True)
            previous_session = (day, session)
            previous_type = None
        if round_.event_type is not previous_type:
            type_title = document.add_paragraph()
            type_title.paragraph_format.keep_with_next = True
            type_title.paragraph_format.space_before = Pt(8)
            type_title.paragraph_format.space_after = Pt(4)
            _set_run_font(type_title.add_run(type_names[round_.event_type]), "黑体", 12, True)
            previous_type = round_.event_type
        event_title = document.add_paragraph()
        # 标题必须与紧随其后的内容同页：有名单时跟“第N组”表头，没有名单时
        # （集体项目、教工项目）跟下一条标题。否则末条标题会被单独留在页尾。
        event_title.paragraph_format.keep_with_next = True
        event_title.paragraph_format.space_before = Pt(8)
        event_title.paragraph_format.space_after = Pt(3)
        count = round_.declared_count if round_.declared_count is not None else len(event_assignments)
        suffix = {
            RoundType.PRELIMINARY: "预赛",
            RoundType.FINAL: "决赛",
            RoundType.TIMED_FINAL: "预决赛",
        }[round_.round_type]
        time_text = _schedule_parts(round_.scheduled_time)[2] if round_.scheduled_time else ""
        count_unit = "队" if round_.event_type is EventType.TEAM else "人"
        run = event_title.add_run(
            f"{sequence}：{round_.group_name}{round_.event_name}{suffix}\t{count}{count_unit}\t共{round_.heat_count}组\t{time_text}"
        )
        _set_run_font(run, "宋体", 11, bold=True)

        if not event_assignments:
            continue
        heat_numbers = sorted({item.heat_no for item in event_assignments})
        for heat_no in heat_numbers:
            group_title = document.add_paragraph(f"第{heat_no}组")
            group_title.paragraph_format.keep_with_next = True
            group_title.paragraph_format.space_after = Pt(2)
            group = [item for item in event_assignments if item.heat_no == heat_no]
            group.sort(key=lambda item: item.lane if item.lane is not None else item.order or 0)
            table = document.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = ["道次" if round_.event_type is EventType.TRACK else "序号", "号码", "姓名", "单位", "成绩"]
            for cell, value in zip(table.rows[0].cells, headers, strict=True):
                _set_cell_text(cell, value, bold=True)
            for assignment in group:
                athlete = athlete_map[assignment.participant_id]
                number = assignment.lane if assignment.lane is not None else assignment.order
                row = table.add_row().cells
                values = [number, athlete.bib, athlete.name, f"{athlete.unit}班", ""]
                for cell, value in zip(row, values, strict=True):
                    _set_cell_text(
                        cell,
                        "" if value is None else str(value),
                        single_line_capacity=6.0,
                        diagnostics=diagnostics,
                        context=f"{round_.group_name}{round_.event_name}第{heat_no}组",
                    )
            for row in table.rows:
                for cell in row.cells:
                    cell.width = Cm(3.048)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _prevent_table_row_split(table)
            table_count += 1

    return RenderStats(len(ordered_rounds), table_count, len(assignments))


def _append_schedule(
    document: Document, rounds: list[EventRound], diagnostics: _LayoutDiagnostics | None = None
) -> None:
    _add_chapter_title(document, "五、竞赛日程")
    type_names = {EventType.TRACK: "径赛", EventType.FIELD: "田赛", EventType.TEAM: "集体项目"}
    round_names = {
        RoundType.PRELIMINARY: "预赛", RoundType.FINAL: "决赛", RoundType.TIMED_FINAL: "预决赛",
    }
    blocks = defaultdict(list)
    for round_ in sorted(rounds, key=lambda item: item.scheduled_time):
        day, session, clock = _schedule_parts(round_.scheduled_time)
        blocks[(day, session)].append((round_, clock))
    # 每个时段内固定按“先径赛 → 再田赛 → 后集体项目”排列，而不是按时间交错。
    type_order = {EventType.TRACK: 0, EventType.FIELD: 1, EventType.TEAM: 2}
    for events in blocks.values():
        events.sort(key=lambda item: (type_order[item[0].event_type], item[1]))
    for block_index, ((day, session), events) in enumerate(blocks.items()):
        heading = document.add_paragraph()
        heading.paragraph_format.page_break_before = block_index > 0
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(6)
        heading.paragraph_format.keep_with_next = True
        _set_run_font(heading.add_run(_session_label(day, session)), "黑体", 14, True)
        for kind in (EventType.TRACK, EventType.FIELD, EventType.TEAM):
            selected = [(event, clock) for event, clock in events if event.event_type is kind]
            if not selected:
                continue
            subheading = document.add_paragraph()
            subheading.paragraph_format.keep_with_next = True
            subheading.paragraph_format.space_before = Pt(6)
            subheading.paragraph_format.space_after = Pt(4)
            _set_run_font(subheading.add_run(type_names[kind]), "黑体", 12, True)
            table = document.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            widths = (2.0, 8.0, 2.0, 2.5, 2.1)
            for column, width in zip(table.columns, widths, strict=True):
                column.width = Cm(width)
            for cell, value in zip(table.rows[0].cells, ("比赛时间", "竞赛项目", "赛次", "人数/队数", "组数"), strict=True):
                _set_cell_text(cell, value, bold=True)
            for event, clock in selected:
                unit = "队" if kind is EventType.TEAM else "人"
                count = "" if event.declared_count is None else f"{event.declared_count}{unit}"
                values = (clock, f"{event.group_name}{event.event_name}", round_names[event.round_type], count, f"{event.heat_count}组")
                for index, (cell, value) in enumerate(zip(table.add_row().cells, values, strict=True)):
                    _set_cell_text(cell, value)
                    if index == 1:
                        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            for row in table.rows:
                for cell, width in zip(row.cells, widths, strict=True):
                    cell.width = Cm(width)
            _prevent_table_row_split(table)


def _schedule_parts(value: str) -> tuple[str, str, str]:
    day = re.search(r"\d{4}-\d{1,2}-\d{1,2}", value)
    clocks = re.findall(r"\d{1,2}\s*[:：]\s*\d{2}", value)
    clock = re.sub(r"\s|：", lambda match: ":" if match.group() == "：" else "", clocks[-1]) if clocks else value
    session = "上午" if "上午" in value else "下午" if "下午" in value else ""
    return day.group() if day else "日期待定", session, clock


def _append_static_chapter(document: Document, text: str, placeholder: str, heading: str) -> None:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        document.add_paragraph(placeholder)
        return
    # 导入的原文通常含章节全称，仅在与当前章名一致时去掉重复标题。
    subject = heading.split("、", 1)[-1]
    first_line = re.sub(r"\s", "", lines[0])
    if first_line.endswith(subject) and ("运动会" in first_line or first_line == subject):
        lines = lines[1:]
    compact = "裁判员" in heading
    opening = "开幕式" in heading
    lines = [part for line in lines for part in re.split(r"\s+(?=[\u4e00-\u9fff]{2,12}组组长[：:])", line)]
    for line in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = Pt(14 if compact else 22 if opening else 20)
        paragraph.paragraph_format.space_after = Pt(1 if compact else 4)
        paragraph.paragraph_format.keep_together = True
        size = 11 if compact else 12
        label = re.match(r"^([^：:]{1,18}[：:])(.*)$", line)
        if label:
            _set_run_font(paragraph.add_run(label[1]), "宋体", size, True)
            _set_run_font(paragraph.add_run(label[2]), "宋体", size)
        else:
            _set_run_font(paragraph.add_run(line), "宋体", size)
        if line.startswith(("备注", "注：", "注:")):
            previous = paragraph._p.getprevious()
            if previous is not None and previous.tag == qn("w:p"):
                Paragraph(previous, document._body).paragraph_format.keep_with_next = True
        if _is_signature_line(line):
            # 落款（校运会组委会 / 2026年10月）靠右，并与其相邻行保持同页。
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = False
            previous = paragraph._p.getprevious()
            if previous is not None and previous.tag == qn("w:p"):
                previous_paragraph = Paragraph(previous, document._body)
                previous_paragraph.paragraph_format.keep_with_next = True
                previous_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _is_signature_line(line: str) -> bool:
    """识别静态章节末尾的落款行，例如“校运会组委会”“2026年10月”。"""
    text = line.strip()
    if not text:
        return False
    if text in {"校运会组委会", "运动会组委会", "组委会"}:
        return True
    return bool(re.fullmatch(r"\d{4}\s*年\s*\d{1,2}\s*月(?:\s*\d{1,2}\s*日)?", text))


def _append_team_rosters(
    document: Document, athletes: list[Athlete], diagnostics: _LayoutDiagnostics | None = None
) -> None:
    _add_chapter_title(document, "六、运动员代表队名单")
    grouped: dict[tuple[str, str], list[Athlete]] = defaultdict(list)
    for athlete in athletes:
        grouped[(athlete.grade, athlete.class_name)].append(athlete)
    for (grade, class_name), members in sorted(grouped.items(), key=_unit_sort_key):
        heading = document.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        _set_run_font(heading.add_run(f"{grade}({class_name})班"), "宋体", 12, bold=True)
        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        for index, member in enumerate(sorted(members, key=lambda item: item.bib)):
            row = table.rows[index // 3].cells if index % 3 else (
                table.rows[0].cells if index == 0 else table.add_row().cells
            )
            offset = (index % 3) * 2
            _set_cell_text(row[offset], member.bib)
            _set_cell_text(
                row[offset + 1], member.name, single_line_capacity=4.5,
                diagnostics=diagnostics, context=f"{grade}{class_name}班名单",
            )
        _prevent_table_row_split(table)
        # 常规班级名单整块排版，避免最后几名学生单独落到下一页。
        # 超大班级允许跨页，不能把比整页还长的表强行绑定。
        if len(table.rows) <= 12:
            for row in table.rows[:-1]:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True


def _unit_sort_key(item: tuple[tuple[str, str], list[Athlete]]) -> tuple[int, int, str]:
    (grade, class_name), _members = item
    grade_order = {"初一": 1, "初二": 2, "初三": 3, "高一": 4, "高二": 5, "高三": 6}
    try:
        class_number = int(class_name)
    except ValueError:
        class_number = 999
    return grade_order.get(grade, 999), class_number, class_name


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    _configure_normal_style(document)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        properties = document.styles[name]._element.pPr
        if properties is not None:
            for border in list(properties.findall(qn("w:pBdr"))):
                properties.remove(border)
    # 首页页眉页脚是否单独设置由 render_full_booklet 按有无封面决定：
    # 没有封面时必须为 False，否则第 1 页（目录页）不会显示页码。
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(footer.add_run("— "), "宋体", 9)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    _set_run_font(footer.add_run(" —"), "宋体", 9)


def _set_blank_run_font(paragraph) -> None:
    """给空白页脚段落设置字体，保证 Word 不套用异常默认字形。"""
    _set_run_font(paragraph.add_run(""), "宋体", 9)


def _restart_page_numbering(document: Document, start: int = 1) -> None:
    """把页码从指定值重新开始，使封面不占页码（目录页显示为第 1 页）。

    注意：w:pgNumType 是“节”级属性，必须作用在封面之后的节上，否则封面仍算第 1 页。
    """
    section = document.sections[-1]
    properties = section._sectPr
    for existing in properties.findall(qn("w:pgNumType")):
        properties.remove(existing)
    page_numbering = OxmlElement("w:pgNumType")
    page_numbering.set(qn("w:start"), str(start))
    cols = properties.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(page_numbering)
    else:
        properties.append(page_numbering)


def _add_page_break(document: Document) -> None:
    document.add_page_break()


def _add_chapter_title(document: Document, text: str, size: float = 22, page_break: bool = True) -> None:
    # 标题自身分页，避免单独的分页符段落撑出空白页。
    # page_break=False 只用于文档首个章节（无封面时的目录），否则会多出空白页。
    if document.paragraphs:
        previous = document.paragraphs[-1]
        if not previous.text.strip() and previous._p.xpath('./w:r/w:br[@w:type="page"]'):
            document._body._body.remove(previous._p)
            page_break = True
    title = document.add_paragraph(text, style="Heading 1")
    title.paragraph_format.page_break_before = page_break
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.keep_with_next = True
    for run in title.runs:
        _set_run_font(run, "微软雅黑", size, bold=True)


def _append_multiline_text(
    document: Document,
    text: str,
    placeholder: str,
    diagnostics: _LayoutDiagnostics | None = None,
    table_context: str = "静态章节表格",
) -> None:
    content = text.strip()
    if not content:
        paragraph = document.add_paragraph(placeholder)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return
    lines = content.splitlines()
    rows = [line.split("\t") for line in lines]
    if rows and len(rows[0]) > 1 and all(len(row) == len(rows[0]) for row in rows):
        table = document.add_table(rows=1, cols=len(rows[0]))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, value in zip(table.rows[0].cells, rows[0], strict=True):
            _set_cell_text(cell, value, bold=True)
        for values in rows[1:]:
            for cell, value in zip(table.add_row().cells, values, strict=True):
                _set_cell_text(
                    cell,
                    value,
                    single_line_capacity=6.0,
                    diagnostics=diagnostics,
                    context=f"{table_context}第{len(table.rows)}行",
                )
        _prevent_table_row_split(table)
        return
    for line in lines:
        paragraph = document.add_paragraph(line)
        paragraph.paragraph_format.space_after = Pt(3)


def _configure_normal_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.15
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _set_run_font(run, name: str, size: float, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _set_cell_text(
    cell,
    value: str,
    bold: bool = False,
    single_line_capacity: float | None = None,
    diagnostics: _LayoutDiagnostics | None = None,
    context: str = "表格",
) -> None:
    if single_line_capacity is not None:
        value = normalize_single_line_text(value)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0
    # 不受 Word 默认文档网格的20磅行距约束，否则紧凑表格仍会被撑高。
    snap_to_grid = OxmlElement("w:snapToGrid")
    snap_to_grid.set(qn("w:val"), "0")
    paragraph._p.get_or_add_pPr().append(snap_to_grid)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    run = paragraph.add_run(value)
    font_size = 10.5
    if single_line_capacity is not None:
        fit = fit_single_line_font(value, single_line_capacity)
        font_size = fit.font_size
        if fit.adjusted and diagnostics is not None:
            diagnostics.font_adjustment_count += 1
        if fit.fits:
            prevent_cell_text_wrap(cell)
        elif diagnostics is not None:
            diagnostics.warnings.append(f"{context}：“{value}”过长，已缩至{font_size:g}磅但仍可能换行")
    _set_run_font(run, "宋体", font_size, bold)


def _prevent_table_row_split(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    # 代表队名单没有表头，不把第一排运动员误标为重复表头。
    if table.cell(0, 0).text in {"比赛时间", "道次", "序号", "项目"}:
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F0F0F0")
            cell._tc.get_or_add_tcPr().append(shading)
        if table.cell(0, 0).text in {"道次", "序号"} and len(table.rows) <= 12:
            for row in table.rows[:-1]:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        properties.append(cant_split)
